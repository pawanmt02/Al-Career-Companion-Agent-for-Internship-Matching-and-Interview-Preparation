import os
import re
import tempfile
from datetime import datetime, timedelta
from typing import Any, Dict, List, Optional

from fastapi import FastAPI, File, HTTPException, UploadFile, Depends, status
from fastapi.security import OAuth2PasswordBearer, OAuth2PasswordRequestForm
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, EmailStr, Field
import numpy as np
from sklearn.metrics.pairwise import cosine_similarity
from sentence_transformers import SentenceTransformer

# --- Database & Auth Imports ---
from sqlalchemy import create_engine, Column, Integer, String, Text, ForeignKey, JSON
from sqlalchemy.orm import declarative_base, sessionmaker, Session
from passlib.context import CryptContext
import bcrypt
from jose import JWTError, jwt

# Try importing PDF and Word readers
try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    import docx
except ImportError:
    docx = None

# ==============================================================================
# 1. APP INITIALIZATION & SECURITY CONFIG
# ==============================================================================
app = FastAPI(
    title="AI Resume Agent",
    version="5.0.0",
    description="Secure User Management, AI Resume Parsing, Profile Auto-Fill & Cover Letter Generation."
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://127.0.0.1:5173"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# JWT Config
SECRET_KEY = "super_secret_assignment_key"
ALGORITHM = "HS256"
ACCESS_TOKEN_EXPIRE_MINUTES = 60

pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")
oauth2_scheme = OAuth2PasswordBearer(tokenUrl="auth/login")

# Database Config (SQLite)
SQLALCHEMY_DATABASE_URL = "sqlite:///./airesume.db"
engine = create_engine(SQLALCHEMY_DATABASE_URL, connect_args={"check_same_thread": False})
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
Base = declarative_base()


# ==============================================================================
class DBUser(Base):
    __tablename__ = "users"
    id              = Column(Integer, primary_key=True, index=True)
    full_name       = Column(String, index=True)
    email           = Column(String, unique=True, index=True)
    hashed_password = Column(String)

class DBResume(Base):
    __tablename__ = "resumes"
    id             = Column(Integer, primary_key=True, index=True)
    user_id        = Column(Integer, ForeignKey("users.id"))
    filename       = Column(String)
    extracted_data = Column(JSON)

# NEW: User profile table — auto-populated from resume
class DBUserProfile(Base):
    __tablename__ = "user_profiles"
    id             = Column(Integer, primary_key=True, index=True)
    user_id        = Column(Integer, ForeignKey("users.id"), unique=True)
    phone          = Column(String, nullable=True)
    summary        = Column(Text, nullable=True)
    skills         = Column(JSON, nullable=True)
    education      = Column(Text, nullable=True)
    experience     = Column(Text, nullable=True)
    projects       = Column(Text, nullable=True)
    certifications = Column(Text, nullable=True)
    resume_email   = Column(String, nullable=True)  # email extracted from resume

# NEW: Applications table — tracks which internships users have applied to
class DBApplication(Base):
    __tablename__ = "applications"
    id              = Column(Integer, primary_key=True, index=True)
    user_id         = Column(Integer, ForeignKey("users.id"))
    internship_id   = Column(String, nullable=False)
    internship_title = Column(String, nullable=False)
    company         = Column(String, nullable=False)
    work_mode       = Column(String, nullable=True)
    duration        = Column(String, nullable=True)
    location        = Column(String, nullable=True)
    status          = Column(String, default="Applied")  # Applied, Withdrawn
    applied_at      = Column(String, nullable=False)     # ISO timestamp

# NEW: Document uploads for Q&A
class DBDocument(Base):
    __tablename__ = "documents"
    id             = Column(Integer, primary_key=True, index=True)
    user_id        = Column(Integer, ForeignKey("users.id"))
    filename       = Column(String, nullable=False)
    extracted_text = Column(Text, nullable=False)
    chunks         = Column(JSON, nullable=True)   # list of text chunks
    uploaded_at    = Column(String, nullable=False)

Base.metadata.create_all(bind=engine)

def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()

# ==============================================================================
# 3. SCHEMAS & AUTH UTILS
# ==============================================================================
class UserCreate(BaseModel):
    full_name: str
    email: EmailStr
    password: str = Field(max_length=72)

class UserResponse(BaseModel):
    id: int
    email: EmailStr
    full_name: str
    class Config:
        from_attributes = True

class UserProfileDataResponse(BaseModel):
    phone: Optional[str]          = None
    summary: Optional[str]        = None
    skills: Optional[List[str]]   = None
    education: Optional[str]      = None
    experience: Optional[str]     = None
    projects: Optional[str]       = None
    certifications: Optional[str] = None
    resume_email: Optional[str]   = None
    class Config:
        from_attributes = True

class PasswordChange(BaseModel):
    old_password: str
    new_password: str = Field(max_length=72)

class Token(BaseModel):
    access_token: str
    token_type: str

def verify_password(plain_password, hashed_password):
    return bcrypt.checkpw(plain_password[:72].encode("utf-8"), hashed_password.encode("utf-8"))

def get_password_hash(password):
    return bcrypt.hashpw(password[:72].encode("utf-8"), bcrypt.gensalt()).decode("utf-8")

def create_access_token(data: dict, expires_delta: Optional[timedelta] = None):
    to_encode = data.copy()
    expire = datetime.utcnow() + (expires_delta if expires_delta else timedelta(minutes=60))
    to_encode.update({"exp": expire})
    return jwt.encode(to_encode, SECRET_KEY, algorithm=ALGORITHM)

async def get_current_user(token: str = Depends(oauth2_scheme), db: Session = Depends(get_db)):
    credentials_exception = HTTPException(
        status_code=status.HTTP_401_UNAUTHORIZED,
        detail="Could not validate credentials",
        headers={"WWW-Authenticate": "Bearer"},
    )
    try:
        payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
        email: str = payload.get("sub")
        if email is None:
            raise credentials_exception
    except JWTError:
        raise credentials_exception
    user = db.query(DBUser).filter(DBUser.email == email).first()
    if user is None:
        raise credentials_exception
    return user

# ==============================================================================
# 4. USER MANAGEMENT APIs
# ==============================================================================
@app.post("/auth/register", response_model=UserResponse, status_code=status.HTTP_201_CREATED, tags=["Authentication"])
def register_user(user: UserCreate, db: Session = Depends(get_db)):
    db_user = db.query(DBUser).filter(DBUser.email == user.email).first()
    if db_user:
        raise HTTPException(status_code=400, detail="Email already registered")
    hashed_password = get_password_hash(user.password)
    new_user = DBUser(email=user.email, hashed_password=hashed_password, full_name=user.full_name)
    db.add(new_user)
    db.commit()
    db.refresh(new_user)
    return new_user

@app.post("/auth/login", response_model=Token, tags=["Authentication"])
def login_user(form_data: OAuth2PasswordRequestForm = Depends(), db: Session = Depends(get_db)):
    user = db.query(DBUser).filter(DBUser.email == form_data.username).first()
    if not user or not verify_password(form_data.password, user.hashed_password):
        raise HTTPException(status_code=401, detail="Incorrect email or password")
    access_token = create_access_token(data={"sub": user.email})
    return {"access_token": access_token, "token_type": "bearer"}

@app.put("/auth/change-password", tags=["Authentication"])
def change_password(passwords: PasswordChange, db: Session = Depends(get_db), current_user: DBUser = Depends(get_current_user)):
    if not verify_password(passwords.old_password, current_user.hashed_password):
        raise HTTPException(status_code=400, detail="Incorrect old password")
    current_user.hashed_password = get_password_hash(passwords.new_password)
    db.commit()
    return {"message": "Password updated successfully"}

@app.get("/users/profile", response_model=UserResponse, tags=["User Management"])
def get_user_profile(current_user: DBUser = Depends(get_current_user)):
    return current_user

# NEW: Get resume-extracted profile data
@app.get("/users/profile/resume-data", response_model=UserProfileDataResponse, tags=["User Management"])
def get_resume_profile_data(db: Session = Depends(get_db), current_user: DBUser = Depends(get_current_user)):
    profile = db.query(DBUserProfile).filter(DBUserProfile.user_id == current_user.id).first()
    if not profile:
        raise HTTPException(status_code=404, detail="No resume data found. Please upload a resume first.")
    return profile

@app.delete("/users/profile", status_code=status.HTTP_204_NO_CONTENT, tags=["User Management"])
def delete_user_account(db: Session = Depends(get_db), current_user: DBUser = Depends(get_current_user)):
    # Also delete profile data, applications and documents
    db.query(DBUserProfile).filter(DBUserProfile.user_id == current_user.id).delete()
    db.query(DBApplication).filter(DBApplication.user_id == current_user.id).delete()
    db.query(DBDocument).filter(DBDocument.user_id == current_user.id).delete()
    db.delete(current_user)
    db.commit()

# ==============================================================================
# 4b. INTERNSHIP APIs
# ==============================================================================
INTERNSHIP_DATABASE = [
    {
        "id": "INT-001", "title": "Generative AI Research Intern", "company": "NexusAI Labs",
        "work_mode": "Remote", "duration": "6 Months",
        "required_skills": ["Python", "PyTorch", "Hugging Face", "LLMs", "RAG", "LangChain"],
        "description": "Work on cutting-edge generative AI research, fine-tuning large language models and building RAG pipelines.",
        "location": "Remote (India)"
    },
    {
        "id": "INT-002", "title": "Backend Engineering Intern", "company": "CloudScale Systems",
        "work_mode": "Hybrid", "duration": "3 Months",
        "required_skills": ["Python", "FastAPI", "PostgreSQL", "REST APIs", "Docker"],
        "description": "Build and maintain scalable REST APIs for cloud-based products used by thousands of users.",
        "location": "Bengaluru, India"
    },
    {
        "id": "INT-003", "title": "Data Science & Analytics Intern", "company": "FinMetrics Global",
        "work_mode": "On-site", "duration": "6 Months",
        "required_skills": ["Python", "Pandas", "SQL", "Scikit-Learn", "Data Visualization"],
        "description": "Analyse financial datasets, build predictive models and create interactive dashboards for business intelligence.",
        "location": "Mumbai, India"
    },
    {
        "id": "INT-004", "title": "Full-Stack Web Developer Intern", "company": "DevCraft Studios",
        "work_mode": "Hybrid", "duration": "4 Months",
        "required_skills": ["React.js", "Node.js", "JavaScript", "HTML", "CSS", "REST APIs"],
        "description": "Build responsive, full-stack web applications for clients in e-commerce, edtech, and SaaS sectors.",
        "location": "Pune, India"
    },
    {
        "id": "INT-005", "title": "Machine Learning Engineer Intern", "company": "VisionTech AI",
        "work_mode": "Remote", "duration": "6 Months",
        "required_skills": ["Python", "Scikit-Learn", "PyTorch", "Pandas", "SQL", "Data Visualization"],
        "description": "Design, train and deploy computer vision and NLP models to production pipelines.",
        "location": "Remote (India)"
    },
    {
        "id": "INT-006", "title": "Cloud Infrastructure Intern", "company": "AeroCloud Solutions",
        "work_mode": "Hybrid", "duration": "3 Months",
        "required_skills": ["AWS", "Docker", "Python", "REST APIs", "PostgreSQL"],
        "description": "Assist in designing and deploying cloud-native microservices on AWS using Docker and Kubernetes.",
        "location": "Hyderabad, India"
    },
    {
        "id": "INT-007", "title": "Frontend React Developer Intern", "company": "PixelForge Labs",
        "work_mode": "Remote", "duration": "3 Months",
        "required_skills": ["React.js", "TypeScript", "JavaScript", "HTML", "CSS"],
        "description": "Create pixel-perfect, animated UI components for a suite of B2B SaaS products used globally.",
        "location": "Remote (India)"
    },
    {
        "id": "INT-008", "title": "Cybersecurity Analyst Intern", "company": "SecureNet India",
        "work_mode": "On-site", "duration": "6 Months",
        "required_skills": ["Python", "SQL", "REST APIs", "Docker"],
        "description": "Monitor network threats, perform vulnerability assessments and assist in incident response for enterprise clients.",
        "location": "Delhi, India"
    },
    {
        "id": "INT-009", "title": "Data Engineering Intern", "company": "StreamBase Technologies",
        "work_mode": "Hybrid", "duration": "6 Months",
        "required_skills": ["Python", "SQL", "PostgreSQL", "Pandas", "Docker", "AWS"],
        "description": "Build and optimise ETL pipelines that process millions of events per day using Apache Spark and Kafka.",
        "location": "Bengaluru, India"
    },
    {
        "id": "INT-010", "title": "NLP & Chatbot Development Intern", "company": "LinguaBot AI",
        "work_mode": "Remote", "duration": "4 Months",
        "required_skills": ["Python", "LangChain", "LLMs", "Hugging Face", "REST APIs", "FastAPI"],
        "description": "Build intelligent conversational AI chatbots for customer support using LangChain and fine-tuned LLMs.",
        "location": "Remote (India)"
    },
    {
        "id": "INT-011", "title": "Business Intelligence Intern", "company": "InsightFlow Analytics",
        "work_mode": "On-site", "duration": "3 Months",
        "required_skills": ["SQL", "Power BI", "Python", "Pandas", "Data Visualization", "Excel"],
        "description": "Design BI dashboards, write complex SQL queries and present data-driven insights to business stakeholders.",
        "location": "Chennai, India"
    },
    {
        "id": "INT-012", "title": "Android App Development Intern", "company": "MobileFirst Labs",
        "work_mode": "Hybrid", "duration": "4 Months",
        "required_skills": ["Java", "Python", "REST APIs", "SQL"],
        "description": "Develop and publish Android applications, integrate backend APIs and improve app performance.",
        "location": "Gurugram, India"
    },
    {
        "id": "INT-013", "title": "DevOps & Automation Intern", "company": "PipelineX Corp",
        "work_mode": "Remote", "duration": "6 Months",
        "required_skills": ["Docker", "AWS", "Python", "REST APIs", "PostgreSQL"],
        "description": "Automate CI/CD pipelines, manage cloud infrastructure and improve deployment reliability using Docker and AWS.",
        "location": "Remote (India)"
    },
    {
        "id": "INT-014", "title": "Streamlit & Data App Intern", "company": "DataViz Studio",
        "work_mode": "Remote", "duration": "3 Months",
        "required_skills": ["Python", "Streamlit", "Pandas", "SQL", "Data Visualization"],
        "description": "Build interactive data exploration tools and internal analytics dashboards using Streamlit and Python.",
        "location": "Remote (India)"
    },
    {
        "id": "INT-015", "title": "AI Product Research Intern", "company": "FutureStack Ventures",
        "work_mode": "Hybrid", "duration": "6 Months",
        "required_skills": ["Python", "LLMs", "RAG", "FastAPI", "React.js", "REST APIs"],
        "description": "Research and prototype AI-powered product features including semantic search, recommendation systems and AI assistants.",
        "location": "Bengaluru, India"
    },
]

@app.get("/internships", tags=["Internships"])
def list_internships():
    return {"internships": INTERNSHIP_DATABASE, "total": len(INTERNSHIP_DATABASE)}

@app.get("/internships/{internship_id}", tags=["Internships"])
def get_internship(internship_id: str):
    job = next((j for j in INTERNSHIP_DATABASE if j["id"] == internship_id), None)
    if not job:
        raise HTTPException(status_code=404, detail=f"Internship '{internship_id}' not found")
    return job

# Generate cover letter for an internship (no email)
@app.post("/internships/{internship_id}/cover-letter", tags=["Internships"])
def generate_cover_letter_for_internship(
    internship_id: str,
    db: Session = Depends(get_db),
    current_user: DBUser = Depends(get_current_user)
):
    """Generate a personalised cover letter based on the user's resume profile."""
    job = next((j for j in INTERNSHIP_DATABASE if j["id"] == internship_id), None)
    if not job:
        raise HTTPException(status_code=404, detail="Internship not found")

    profile = db.query(DBUserProfile).filter(DBUserProfile.user_id == current_user.id).first()
    if not profile:
        raise HTTPException(
            status_code=400,
            detail="Please upload your resume first so we can generate a personalised cover letter."
        )

    cover_letter = generate_cover_letter(current_user, profile, job)
    return {
        "status": "success",
        "internship": job["title"],
        "company": job["company"],
        "cover_letter": cover_letter,
    }

# ==============================================================================
# 4c. APPLICATION APIs — Track user applications
# ==============================================================================
@app.post("/internships/{internship_id}/apply", tags=["Applications"])
def apply_to_internship(
    internship_id: str,
    db: Session = Depends(get_db),
    current_user: DBUser = Depends(get_current_user)
):
    """Apply to an internship — stores a record in the applications table."""
    job = next((j for j in INTERNSHIP_DATABASE if j["id"] == internship_id), None)
    if not job:
        raise HTTPException(status_code=404, detail="Internship not found")

    # Prevent duplicate applications
    existing = db.query(DBApplication).filter(
        DBApplication.user_id == current_user.id,
        DBApplication.internship_id == internship_id,
        DBApplication.status == "Applied"
    ).first()
    if existing:
        raise HTTPException(status_code=400, detail="You have already applied to this internship.")

    application = DBApplication(
        user_id=current_user.id,
        internship_id=job["id"],
        internship_title=job["title"],
        company=job["company"],
        work_mode=job.get("work_mode"),
        duration=job.get("duration"),
        location=job.get("location"),
        status="Applied",
        applied_at=datetime.utcnow().isoformat()
    )
    db.add(application)
    db.commit()
    db.refresh(application)
    return {
        "status": "success",
        "message": f"Successfully applied to {job['title']} at {job['company']}",
        "application_id": application.id
    }

@app.get("/applications", tags=["Applications"])
def list_my_applications(
    db: Session = Depends(get_db),
    current_user: DBUser = Depends(get_current_user)
):
    """List all applications for the current user."""
    apps = db.query(DBApplication).filter(
        DBApplication.user_id == current_user.id
    ).order_by(DBApplication.id.desc()).all()
    return {
        "applications": [
            {
                "id": a.id,
                "internship_id": a.internship_id,
                "internship_title": a.internship_title,
                "company": a.company,
                "work_mode": a.work_mode,
                "duration": a.duration,
                "location": a.location,
                "status": a.status,
                "applied_at": a.applied_at,
            }
            for a in apps
        ],
        "total": len(apps)
    }

@app.delete("/applications/{application_id}", tags=["Applications"])
def withdraw_application(
    application_id: int,
    db: Session = Depends(get_db),
    current_user: DBUser = Depends(get_current_user)
):
    """Withdraw (delete) an application."""
    app_record = db.query(DBApplication).filter(
        DBApplication.id == application_id,
        DBApplication.user_id == current_user.id
    ).first()
    if not app_record:
        raise HTTPException(status_code=404, detail="Application not found")
    db.delete(app_record)
    db.commit()
    return {"status": "success", "message": "Application withdrawn successfully"}

# ==============================================================================
# 4d. INTERVIEW QUESTION GENERATOR
# ==============================================================================

# Skill-specific technical question bank (with full answers)
SKILL_QUESTIONS = {
    "Python": [
        {"q": "What are Python decorators and how do they work?", "hint": "Explain @syntax, wrapper functions, and common use cases like @staticmethod.",
         "a": "Decorators are functions that modify the behavior of another function without changing its source code. They use the @decorator syntax and work by wrapping the original function inside a wrapper function. For example, @staticmethod and @classmethod are built-in decorators. Custom decorators take a function as input, define an inner wrapper that adds behavior (like logging, timing, or authentication), and return the wrapper."},
        {"q": "Explain the difference between a list and a tuple in Python.", "hint": "Mutability, performance, use cases.",
         "a": "Lists are mutable (can be modified after creation) while tuples are immutable (cannot be changed). Lists use square brackets [] and tuples use parentheses (). Tuples are faster and use less memory since they're immutable, making them ideal for fixed data like coordinates. Lists are better when you need to add, remove, or modify elements. Tuples can be used as dictionary keys, lists cannot."},
        {"q": "What are Python generators and when would you use them?", "hint": "yield keyword, lazy evaluation, memory efficiency.",
         "a": "Generators are functions that use the 'yield' keyword instead of 'return' to produce a sequence of values lazily — one at a time, on demand. They don't store the entire sequence in memory, making them extremely memory-efficient for large datasets. When you call a generator function, it returns a generator object. Each call to next() resumes execution until the next yield. Use generators when processing large files, infinite sequences, or data pipelines."},
        {"q": "How does Python handle memory management and garbage collection?", "hint": "Reference counting, gc module, circular references.",
         "a": "Python uses automatic memory management with two main mechanisms: (1) Reference counting — each object tracks how many references point to it; when the count drops to zero, memory is freed immediately. (2) Generational garbage collector (gc module) — handles circular references that reference counting can't resolve. Objects are grouped into 3 generations; newer objects are collected more frequently. You can manually trigger collection with gc.collect()."},
        {"q": "What is the GIL in Python and how does it affect multithreading?", "hint": "Global Interpreter Lock, CPU-bound vs I/O-bound tasks.",
         "a": "The Global Interpreter Lock (GIL) is a mutex in CPython that allows only one thread to execute Python bytecode at a time. This means multithreading doesn't provide true parallelism for CPU-bound tasks. However, the GIL is released during I/O operations (file reads, network calls), so multithreading is still beneficial for I/O-bound tasks. For CPU-bound parallelism, use multiprocessing module or alternative interpreters like PyPy."},
    ],
    "JavaScript": [
        {"q": "Explain closures in JavaScript with an example.", "hint": "Lexical scoping, inner function accessing outer variables.",
         "a": "A closure is a function that retains access to variables from its outer (enclosing) scope even after the outer function has returned. Example: function outer() { let count = 0; return function inner() { count++; return count; } } const counter = outer(); counter() returns 1, counter() returns 2. The inner function 'closes over' the count variable. Closures are used for data privacy, factory functions, and callbacks."},
        {"q": "What is the difference between var, let, and const?", "hint": "Hoisting, block scope vs function scope, reassignment.",
         "a": "var is function-scoped and hoisted (initialized as undefined), which can cause bugs. let is block-scoped and hoisted but not initialized (Temporal Dead Zone). const is also block-scoped but cannot be reassigned after initialization (though objects/arrays it references can still be mutated). Best practice: use const by default, let when reassignment is needed, avoid var entirely."},
        {"q": "How does the event loop work in JavaScript?", "hint": "Call stack, callback queue, microtasks, macrotasks.",
         "a": "JavaScript is single-threaded. The event loop continuously checks: (1) Execute all synchronous code on the call stack. (2) Process all microtasks (Promises, queueMicrotask). (3) Process one macrotask (setTimeout, setInterval, I/O). (4) Repeat. When async operations complete (API calls, timers), their callbacks are placed in the appropriate queue. The event loop moves callbacks to the call stack only when it's empty."},
        {"q": "What are Promises and async/await in JavaScript?", "hint": "Asynchronous patterns, chaining, error handling.",
         "a": "Promises represent the eventual result of an asynchronous operation. A Promise can be pending, fulfilled, or rejected. You handle results with .then() and errors with .catch(). async/await is syntactic sugar over Promises that makes async code look synchronous. An async function always returns a Promise. 'await' pauses execution until the Promise resolves. Use try/catch for error handling with async/await."},
        {"q": "Explain prototypal inheritance in JavaScript.", "hint": "Prototype chain, Object.create, __proto__.",
         "a": "In JavaScript, every object has an internal [[Prototype]] link to another object. When you access a property, JS first checks the object itself, then walks up the prototype chain until it finds the property or reaches null. Object.create(proto) creates a new object with proto as its prototype. ES6 classes are syntactic sugar over prototypal inheritance. constructor.prototype defines shared methods for all instances."},
    ],
    "React.js": [
        {"q": "What is the Virtual DOM and how does React use it?", "hint": "Diffing algorithm, reconciliation, performance benefits.",
         "a": "The Virtual DOM is a lightweight JavaScript representation of the actual DOM. When state changes, React creates a new Virtual DOM tree, compares it with the previous one using a 'diffing' algorithm (reconciliation), and calculates the minimum number of actual DOM updates needed. This batch-update approach is faster than directly manipulating the real DOM for every change."},
        {"q": "Explain the difference between useState and useEffect hooks.", "hint": "State management vs side effects, dependency arrays.",
         "a": "useState manages local component state — it returns [value, setter] and triggers re-renders when updated. useEffect handles side effects (API calls, subscriptions, DOM manipulation) that run after render. useEffect takes a callback and optional dependency array: empty [] runs once on mount, [dep] runs when dep changes, no array runs every render. Return a cleanup function for subscriptions."},
        {"q": "What are React keys and why are they important?", "hint": "Efficient list rendering, reconciliation, unique identifiers.",
         "a": "Keys are unique identifiers given to elements in a list so React can efficiently track which items changed, were added, or removed during re-renders. Without keys (or with index as key), React may re-render the entire list. Use stable, unique IDs (like database IDs) as keys, not array indices, especially when the list can be reordered or filtered."},
        {"q": "How do you optimize performance in a React application?", "hint": "React.memo, useMemo, useCallback, code splitting.",
         "a": "Key techniques: (1) React.memo — prevents re-renders if props haven't changed. (2) useMemo — memoizes expensive calculations. (3) useCallback — memoizes function references to prevent child re-renders. (4) Code splitting with React.lazy and Suspense. (5) Virtualization for long lists (react-window). (6) Avoiding inline object/array creation in JSX. (7) Using production builds."},
        {"q": "What is Context API and when would you use it over Redux?", "hint": "Prop drilling, global state, complexity trade-offs.",
         "a": "Context API provides a way to share state across components without passing props through every level (prop drilling). Use createContext, Provider, and useContext. Use Context for simple global state like theme, auth, language. Use Redux for complex state with many actions, middleware needs (async logic, logging), time-travel debugging, or when state updates are frequent and need optimization."},
    ],
    "Node.js": [
        {"q": "How does Node.js handle asynchronous operations?", "hint": "Event-driven, non-blocking I/O, libuv.",
         "a": "Node.js uses an event-driven, non-blocking I/O model built on the libuv library. When an async operation (file read, network request) is initiated, Node delegates it to the system kernel or thread pool, and continues executing other code. When the operation completes, a callback is placed on the event queue and processed by the event loop. This allows Node to handle thousands of concurrent connections with a single thread."},
        {"q": "What is middleware in Express.js?", "hint": "Request pipeline, next(), error handling middleware.",
         "a": "Middleware are functions that execute during the request-response cycle, with access to req, res, and next(). They can modify the request/response, end the cycle, or pass control to the next middleware via next(). Types include application-level (app.use), router-level, error-handling (4 params including err), and third-party (cors, morgan). They execute in the order they're registered."},
        {"q": "Explain the difference between process.nextTick() and setImmediate().", "hint": "Event loop phases, execution priority.",
         "a": "process.nextTick() callbacks execute immediately after the current operation completes, before the event loop continues — they're in a 'nextTick queue' with highest priority. setImmediate() executes in the next iteration of the event loop (check phase). nextTick can starve I/O if used recursively, while setImmediate is safer for yielding to the event loop."},
    ],
    "SQL": [
        {"q": "What is the difference between INNER JOIN, LEFT JOIN, and FULL OUTER JOIN?", "hint": "Matching rows, null handling, Venn diagram analogy.",
         "a": "INNER JOIN returns only rows that have matching values in both tables. LEFT JOIN returns all rows from the left table plus matched rows from the right (NULLs where no match). RIGHT JOIN is the opposite. FULL OUTER JOIN returns all rows from both tables, with NULLs where there's no match on either side. Think of it as a Venn diagram: INNER is the intersection, LEFT includes all of the left circle."},
        {"q": "Explain database normalization and its forms (1NF, 2NF, 3NF).", "hint": "Redundancy reduction, functional dependencies.",
         "a": "Normalization organizes data to reduce redundancy. 1NF: Each column has atomic values, no repeating groups. 2NF: Meets 1NF + all non-key columns fully depend on the entire primary key (eliminates partial dependencies). 3NF: Meets 2NF + no non-key column depends on another non-key column (eliminates transitive dependencies). Over-normalization can hurt read performance, so denormalization is sometimes used."},
        {"q": "What are indexes in SQL and how do they improve performance?", "hint": "B-tree, query optimization, trade-offs with writes.",
         "a": "Indexes are data structures (typically B-trees) that speed up data retrieval by creating a sorted reference to table rows. Like a book's index, they let the database find rows without scanning the entire table. However, indexes slow down INSERT/UPDATE/DELETE operations since the index must also be updated. Use indexes on frequently queried columns, foreign keys, and WHERE/ORDER BY columns. Don't over-index."},
        {"q": "Write a query to find the second highest salary from an employees table.", "hint": "Subquery, LIMIT OFFSET, DENSE_RANK().",
         "a": "Method 1: SELECT MAX(salary) FROM employees WHERE salary < (SELECT MAX(salary) FROM employees); Method 2: SELECT salary FROM employees ORDER BY salary DESC LIMIT 1 OFFSET 1; Method 3 (handles ties): SELECT salary FROM (SELECT salary, DENSE_RANK() OVER (ORDER BY salary DESC) as rnk FROM employees) t WHERE rnk = 2;"},
    ],
    "PostgreSQL": [
        {"q": "What are the advantages of PostgreSQL over MySQL?", "hint": "JSONB, CTEs, advanced indexing, ACID compliance.",
         "a": "PostgreSQL offers: JSONB for efficient JSON storage and querying, CTEs (WITH queries) for complex queries, advanced indexing (GIN, GiST, BRIN), full ACID compliance, window functions, materialized views, custom types, better concurrency with MVCC, and support for geospatial data via PostGIS. MySQL is simpler and faster for read-heavy workloads, but PostgreSQL is more feature-rich and standards-compliant."},
        {"q": "Explain JSONB data type in PostgreSQL.", "hint": "Binary JSON, indexing, querying nested data.",
         "a": "JSONB stores JSON in binary format, which is faster to process than plain JSON text. You can index JSONB columns with GIN indexes for fast lookups. Query nested data with operators: -> (get object), ->> (get text), @> (contains), ? (key exists). Example: SELECT data->>'name' FROM users WHERE data @> '{\"role\": \"admin\"}'. JSONB also supports partial updates and is ideal for semi-structured data."},
    ],
    "FastAPI": [
        {"q": "What makes FastAPI faster than Flask or Django?", "hint": "ASGI, async/await, Pydantic validation, auto-docs.",
         "a": "FastAPI is built on ASGI (Starlette) supporting native async/await for non-blocking I/O, making it significantly faster for concurrent requests. It uses Pydantic for automatic request validation and serialization with zero boilerplate. It auto-generates OpenAPI docs (Swagger/ReDoc). Type hints provide editor support and catch errors early. Benchmarks show it's comparable to Node.js and Go in performance."},
        {"q": "How does dependency injection work in FastAPI?", "hint": "Depends(), reusable dependencies, request lifecycle.",
         "a": "FastAPI uses the Depends() function for dependency injection. You define a function (sync or async) that provides a resource (database session, current user, config), then declare it as a parameter with Depends(). FastAPI resolves the dependency tree automatically. Dependencies can depend on other dependencies, creating a chain. They're great for auth, DB sessions, and shared logic. Use yield for cleanup (like closing DB connections)."},
        {"q": "Explain how Pydantic models work for request validation.", "hint": "Type hints, automatic validation, serialization.",
         "a": "Pydantic models are Python classes that inherit from BaseModel and use type annotations to define fields. When used as FastAPI request bodies, incoming JSON is automatically validated against the model — wrong types, missing required fields, or invalid values return clear 422 error responses. Pydantic also handles serialization (model.dict()), default values, custom validators (@validator), and nested models."},
    ],
    "Django": [
        {"q": "Explain the MTV (Model-Template-View) architecture in Django.", "hint": "ORM models, template rendering, view logic.",
         "a": "MTV stands for Model-Template-View: Model defines data structure using Django ORM (maps to database tables). Template handles presentation (HTML with Django template language for dynamic content). View contains business logic — receives HTTP requests, interacts with Models, and returns responses using Templates. It's similar to MVC but Django's 'View' is like a Controller, and 'Template' is like a View."},
        {"q": "What is Django ORM and how does it handle database queries?", "hint": "QuerySets, lazy evaluation, migrations.",
         "a": "Django ORM maps Python classes to database tables, allowing you to interact with databases using Python instead of SQL. QuerySets are lazy — they don't hit the database until evaluated (iteration, slicing, list()). Common operations: Model.objects.filter(), .exclude(), .annotate(), .aggregate(). Chaining is efficient. Migrations auto-generate SQL from model changes. The ORM supports multiple databases and raw SQL when needed."},
    ],
    "Docker": [
        {"q": "What is the difference between a Docker image and a container?", "hint": "Blueprint vs running instance, layers, immutability.",
         "a": "A Docker image is an immutable blueprint/template containing the application code, runtime, libraries, and dependencies — built in read-only layers. A container is a running instance of an image — it adds a writable layer on top. You can create multiple containers from one image. Think of an image as a class and a container as an object/instance. Images are stored in registries, containers run on Docker engine."},
        {"q": "How do you optimize a Dockerfile for smaller image sizes?", "hint": "Multi-stage builds, alpine base, layer caching.",
         "a": "Key techniques: (1) Use multi-stage builds — compile in one stage, copy only the binary to a minimal final stage. (2) Use alpine-based images (python:3.11-alpine) instead of full OS images. (3) Combine RUN commands with && to reduce layers. (4) Use .dockerignore to exclude unnecessary files. (5) Order instructions from least to most frequently changed for layer caching. (6) Remove package manager caches in the same RUN step."},
        {"q": "Explain Docker Compose and its use cases.", "hint": "Multi-container apps, networking, volumes, YAML config.",
         "a": "Docker Compose is a tool for defining and running multi-container applications using a docker-compose.yml file. You define services (containers), networks, and volumes declaratively. It handles container orchestration locally — starting, stopping, and rebuilding services together. Use cases: running app + database + redis together, development environments, CI/CD testing. Commands: docker compose up, down, build, logs."},
    ],
    "AWS": [
        {"q": "Explain the difference between EC2, Lambda, and ECS.", "hint": "VMs, serverless, container orchestration.",
         "a": "EC2 provides virtual machines (instances) with full OS control — you manage scaling, patching, and availability. Lambda is serverless — you upload code, AWS runs it on demand, scales automatically, and you pay per invocation (no idle cost). ECS (Elastic Container Service) orchestrates Docker containers — manages deployment, scaling, and networking of containerized apps. Choose based on control vs convenience trade-off."},
        {"q": "What is an S3 bucket and what are its storage classes?", "hint": "Object storage, Standard, IA, Glacier, lifecycle policies.",
         "a": "S3 (Simple Storage Service) is object storage for files of any size. Buckets are containers for objects. Storage classes optimize cost vs access speed: Standard (frequent access), Intelligent-Tiering (auto-moves between tiers), Standard-IA (infrequent access, lower cost), One Zone-IA (single AZ), Glacier Instant/Flexible/Deep Archive (archival, cheapest, minutes to hours retrieval). Lifecycle policies auto-transition objects between classes."},
        {"q": "How does IAM work in AWS?", "hint": "Users, roles, policies, principle of least privilege.",
         "a": "IAM (Identity and Access Management) controls who can access what in AWS. Users are individual identities, Groups organize users, Roles are temporary identities assumed by services/users. Policies are JSON documents defining permissions (Allow/Deny actions on resources). Follow the Principle of Least Privilege — grant only the minimum permissions needed. Use MFA for security. Roles are preferred over long-term access keys."},
    ],
    "PyTorch": [
        {"q": "What is autograd in PyTorch and how does it work?", "hint": "Computational graph, backward(), gradient tracking.",
         "a": "Autograd is PyTorch's automatic differentiation engine. When you create tensors with requires_grad=True, PyTorch builds a dynamic computational graph tracking all operations. Calling .backward() on the output computes gradients for all tensors in the graph via backpropagation. Gradients are stored in tensor.grad. This dynamic graph (define-by-run) is rebuilt each forward pass, making debugging and dynamic architectures easy."},
        {"q": "Explain the difference between torch.Tensor and torch.nn.Parameter.", "hint": "Gradient tracking, model parameters, optimization.",
         "a": "torch.Tensor is the base data structure for all tensor operations. torch.nn.Parameter is a Tensor subclass that, when assigned as a module attribute, is automatically registered as a model parameter — it appears in model.parameters() and is updated by optimizers during training. Parameters have requires_grad=True by default. Use Parameter for learnable weights; use regular Tensors for fixed data like input."},
        {"q": "How do you handle overfitting in a PyTorch model?", "hint": "Dropout, regularization, early stopping, data augmentation.",
         "a": "Strategies: (1) Dropout (nn.Dropout) — randomly zeros neurons during training. (2) L1/L2 regularization via optimizer weight_decay parameter. (3) Early stopping — monitor validation loss and stop when it stops improving. (4) Data augmentation (torchvision.transforms) — artificially expand training data. (5) Reduce model complexity (fewer layers/neurons). (6) Batch normalization. (7) Cross-validation. (8) Increase training data."},
    ],
    "Scikit-Learn": [
        {"q": "What is the difference between classification and regression?", "hint": "Discrete vs continuous output, metrics, algorithms.",
         "a": "Classification predicts discrete labels/categories (spam vs not spam, cat vs dog) using algorithms like Logistic Regression, SVM, Random Forest, and metrics like accuracy, precision, recall, F1-score. Regression predicts continuous numerical values (house price, temperature) using Linear Regression, Decision Trees, and metrics like MSE, RMSE, R-squared. The choice depends on whether your target variable is categorical or numerical."},
        {"q": "Explain cross-validation and why it is important.", "hint": "K-fold, overfitting prevention, generalization.",
         "a": "Cross-validation splits data into K folds, trains on K-1 folds and validates on the remaining one, rotating through all folds. This gives K performance scores averaged for a robust estimate. It prevents overfitting by ensuring the model generalizes to unseen data, not just the training set. Common types: K-Fold (typically K=5 or 10), Stratified K-Fold (maintains class distribution), Leave-One-Out. Use cross_val_score() in sklearn."},
        {"q": "What is the bias-variance trade-off?", "hint": "Underfitting vs overfitting, model complexity.",
         "a": "Bias is error from oversimplified models (underfitting) — the model misses patterns. Variance is error from overly complex models (overfitting) — the model memorizes noise. High bias = too simple, high variance = too complex. The trade-off: increasing complexity reduces bias but increases variance, and vice versa. The goal is finding the sweet spot that minimizes total error. Techniques: regularization, cross-validation, ensemble methods."},
    ],
    "LLMs": [
        {"q": "What is a Large Language Model and how does it generate text?", "hint": "Transformer architecture, next token prediction, training data.",
         "a": "An LLM is a deep learning model (typically Transformer-based) trained on massive text datasets to understand and generate human language. It works by predicting the next token (word/subword) given the previous context. During training, it learns patterns, grammar, facts, and reasoning from billions of tokens. At inference, it generates text autoregressively — predicting one token at a time, feeding each prediction back as input. Examples: GPT-4, Gemini, LLaMA."},
        {"q": "Explain the difference between fine-tuning and prompt engineering.", "hint": "Weight updates vs input design, cost, flexibility.",
         "a": "Fine-tuning updates the model's weights on domain-specific data — it requires labeled data, GPU compute, and time, but produces highly specialized models. Prompt engineering crafts the input text to guide the model's output without changing weights — it's faster, cheaper, and more flexible but limited by the model's existing knowledge. Fine-tuning is better for specific tasks (medical, legal), prompt engineering for general applications."},
        {"q": "What are hallucinations in LLMs and how can you mitigate them?", "hint": "Grounding, RAG, temperature, fact-checking.",
         "a": "Hallucinations occur when LLMs generate confident but factually incorrect or fabricated information. Mitigation strategies: (1) RAG — retrieve real documents to ground responses in facts. (2) Lower temperature — reduces randomness in generation. (3) Explicit instructions to say 'I don't know'. (4) Chain-of-thought prompting for reasoning. (5) Fact-checking pipelines. (6) Fine-tuning on verified data. (7) Human-in-the-loop review for critical applications."},
    ],
    "RAG": [
        {"q": "What is Retrieval-Augmented Generation and why is it useful?", "hint": "Combine retrieval with generation, reduce hallucination, domain knowledge.",
         "a": "RAG combines information retrieval with text generation. Instead of relying solely on the LLM's training data, RAG first retrieves relevant documents from a knowledge base, then provides them as context to the LLM for generating answers. Benefits: reduces hallucinations by grounding in real data, enables access to private/updated knowledge without retraining, cost-effective compared to fine-tuning, and provides source citations."},
        {"q": "Explain the RAG pipeline: embedding, indexing, retrieval, generation.", "hint": "Vector stores, similarity search, context injection.",
         "a": "The RAG pipeline: (1) Embedding — convert documents into vector representations using models like sentence-transformers. (2) Indexing — store embeddings in a vector database (FAISS, Pinecone, ChromaDB) for efficient search. (3) Retrieval — when a query comes in, embed it and find the most similar document chunks via cosine similarity. (4) Generation — inject retrieved chunks into the LLM prompt as context, and generate an answer grounded in that context."},
        {"q": "What are vector embeddings and how are they used in RAG?", "hint": "Semantic representation, cosine similarity, sentence transformers.",
         "a": "Vector embeddings are dense numerical representations of text that capture semantic meaning — similar texts have similar vectors. Models like sentence-transformers (all-MiniLM-L6-v2) convert text into fixed-size vectors (e.g., 384 dimensions). In RAG, documents are chunked and embedded into a vector store. Queries are also embedded, and cosine similarity finds the most semantically relevant chunks, even if they don't share exact keywords."},
    ],
    "LangChain": [
        {"q": "What is LangChain and what problems does it solve?", "hint": "LLM orchestration, chains, agents, memory.",
         "a": "LangChain is a framework for building applications powered by LLMs. It solves the complexity of chaining multiple LLM calls, managing prompts, connecting to external tools and data sources, and maintaining conversation memory. Key components: Chains (sequential LLM operations), Agents (LLMs that decide which tools to use), Memory (conversation history), and Retrievers (RAG integration). It supports multiple LLM providers."},
        {"q": "Explain the concept of Chains and Agents in LangChain.", "hint": "Sequential processing, tool-using agents, decision making.",
         "a": "Chains are predefined sequences of operations — e.g., LLMChain takes a prompt template and LLM, SequentialChain pipes output of one chain into another. They follow a fixed flow. Agents are more dynamic — they use an LLM to decide which tools to call and in what order based on the user's input. Tools can be search engines, calculators, databases, APIs. Agents reason step-by-step (ReAct pattern) to solve complex tasks autonomously."},
    ],
    "Hugging Face": [
        {"q": "What is the Hugging Face Transformers library?", "hint": "Pre-trained models, pipelines, model hub.",
         "a": "Hugging Face Transformers is an open-source library providing thousands of pre-trained models for NLP, computer vision, and audio tasks. It offers a simple pipeline() API for quick inference, and fine-grained control with AutoModel/AutoTokenizer classes. The Model Hub hosts 200K+ models. Key features: model cards, datasets library, Trainer API for fine-tuning, and support for PyTorch, TensorFlow, and JAX backends."},
        {"q": "How do you fine-tune a pre-trained model using Hugging Face?", "hint": "Trainer API, datasets, tokenization, evaluation.",
         "a": "Steps: (1) Load a pre-trained model and tokenizer (AutoModelForSequenceClassification, AutoTokenizer). (2) Prepare your dataset using the datasets library. (3) Tokenize data using tokenizer with padding/truncation. (4) Define TrainingArguments (learning rate, epochs, batch size). (5) Create a Trainer with model, args, train/eval datasets. (6) Call trainer.train(). (7) Evaluate with trainer.evaluate(). (8) Save with model.save_pretrained()."},
    ],
    "Pandas": [
        {"q": "What is the difference between a Series and a DataFrame in Pandas?", "hint": "1D vs 2D, indexing, column operations.",
         "a": "A Series is a 1-dimensional labeled array that can hold any data type — like a single column. A DataFrame is a 2-dimensional labeled table with rows and columns — like a spreadsheet or SQL table. A DataFrame is essentially a collection of Series sharing the same index. You can access a single column from a DataFrame as a Series using df['column_name']. Series support element-wise operations and vectorized functions."},
        {"q": "How do you handle missing data in Pandas?", "hint": "dropna(), fillna(), interpolation, isnull().",
         "a": "Detect missing data: df.isnull().sum(). Remove: df.dropna() drops rows with any NaN, df.dropna(subset=['col']) targets specific columns. Fill: df.fillna(value) with a constant, df.fillna(method='ffill') forward-fills, df['col'].fillna(df['col'].mean()) fills with mean. Interpolation: df.interpolate() for numerical data. Best practice: understand WHY data is missing before choosing a strategy — random vs systematic missingness matters."},
        {"q": "Explain groupby operations in Pandas.", "hint": "Split-apply-combine, aggregation, transformation.",
         "a": "groupby follows the split-apply-combine pattern: (1) Split data into groups based on column values. (2) Apply a function to each group (aggregation, transformation, or filtration). (3) Combine results. Example: df.groupby('city')['salary'].mean() gives average salary per city. Use .agg() for multiple functions: .agg(['mean','sum','count']). Use .transform() to broadcast results back to original shape. Named aggregation: .agg(avg_salary=('salary','mean'))."},
    ],
    "REST APIs": [
        {"q": "What are the main HTTP methods and when do you use each?", "hint": "GET, POST, PUT, PATCH, DELETE — CRUD mapping.",
         "a": "GET — retrieve data (Read), should be idempotent and safe. POST — create a new resource (Create), not idempotent. PUT — update/replace an entire resource (Update), idempotent. PATCH — partially update a resource, may not be idempotent. DELETE — remove a resource (Delete), idempotent. Mapping to CRUD: Create=POST, Read=GET, Update=PUT/PATCH, Delete=DELETE. GET requests should never modify server state."},
        {"q": "What is the difference between authentication and authorization?", "hint": "Identity verification vs permission checking, JWT, OAuth.",
         "a": "Authentication verifies WHO you are (identity) — login with username/password, JWT tokens, OAuth. Authorization determines WHAT you can do (permissions) — role-based access control (admin vs user), resource ownership checks. Authentication happens first, then authorization. Example: Logging in with credentials = authentication. Checking if logged-in user can delete a post = authorization. JWT tokens carry both identity and permission claims."},
        {"q": "Explain REST API best practices for URL design.", "hint": "Nouns not verbs, versioning, status codes, pagination.",
         "a": "Best practices: (1) Use nouns, not verbs: /users not /getUsers. (2) Use plural nouns: /users/123 not /user/123. (3) Version your API: /api/v1/users. (4) Use proper HTTP methods instead of encoding actions in URLs. (5) Use query params for filtering: /users?role=admin. (6) Return proper status codes: 200 OK, 201 Created, 404 Not Found, 422 Validation Error. (7) Implement pagination: ?page=1&limit=20. (8) Use kebab-case for multi-word paths."},
    ],
    "HTML": [
        {"q": "What is semantic HTML and why is it important?", "hint": "header, nav, main, article — accessibility, SEO.",
         "a": "Semantic HTML uses tags that describe their content's meaning: <header>, <nav>, <main>, <article>, <section>, <aside>, <footer> instead of generic <div>. Benefits: (1) Accessibility — screen readers understand page structure. (2) SEO — search engines better index content. (3) Maintainability — code is self-documenting. (4) Browser features — like Reader Mode use semantic structure. Always prefer semantic elements over divs with class names."},
        {"q": "Explain the difference between block and inline elements.", "hint": "Layout behavior, examples, display property.",
         "a": "Block elements take the full width available, start on a new line, and stack vertically. Examples: <div>, <p>, <h1>-<h6>, <section>. Inline elements only take the width of their content, sit side by side, and don't start a new line. Examples: <span>, <a>, <strong>, <img>. You can change behavior with CSS display property: display:block, display:inline, display:inline-block (inline but with width/height control)."},
    ],
    "CSS": [
        {"q": "What is the CSS Box Model?", "hint": "Content, padding, border, margin — box-sizing.",
         "a": "Every HTML element is a box with four layers from inside out: Content (the actual text/image), Padding (space between content and border), Border (the visible border line), Margin (space outside the border, between elements). By default (content-box), width/height only apply to content. With box-sizing:border-box (recommended), width includes padding and border, making layouts more predictable."},
        {"q": "Explain Flexbox and when you would use it.", "hint": "One-dimensional layout, justify-content, align-items.",
         "a": "Flexbox is a one-dimensional layout model for arranging items in a row or column. Set display:flex on the container. Key properties: flex-direction (row/column), justify-content (main axis alignment: center, space-between), align-items (cross axis: center, stretch), flex-wrap (wrapping), gap (spacing). Use Flexbox for: navigation bars, centering content, equal-height columns, responsive card layouts, and any single-direction arrangement."},
        {"q": "What is the difference between CSS Grid and Flexbox?", "hint": "2D vs 1D, use cases, combining both.",
         "a": "Flexbox is one-dimensional — it handles layout in a single direction (row OR column) at a time. CSS Grid is two-dimensional — it handles rows AND columns simultaneously. Use Flexbox for: navbars, centering, single-row/column layouts. Use Grid for: page layouts, dashboards, image galleries, any 2D layout. They work great together — Grid for the overall page structure, Flexbox for component-level arrangement within grid cells."},
    ],
    "Data Visualization": [
        {"q": "What makes an effective data visualization?", "hint": "Clarity, appropriate chart type, color, labeling.",
         "a": "Key principles: (1) Choose the right chart type for your data (bar for comparison, line for trends, scatter for correlation). (2) Keep it simple — remove chart junk (unnecessary gridlines, decorations). (3) Use clear labels, titles, and legends. (4) Use color purposefully — highlight key data, maintain accessibility. (5) Maintain proper scale — don't truncate axes misleadingly. (6) Tell a story — guide the viewer to the insight. (7) Consider your audience's expertise level."},
        {"q": "When would you use a bar chart vs a line chart?", "hint": "Categorical vs time-series, comparison vs trends.",
         "a": "Bar charts are best for comparing discrete categories (sales by region, survey responses by option) — horizontal bars when labels are long. Line charts show trends over continuous/time-series data (stock prices over months, temperature over hours) — they emphasize change and direction. Don't use line charts for categorical data (implies continuity that doesn't exist). Use grouped/stacked bars for multi-series categorical comparisons."},
    ],
    "Power BI": [
        {"q": "What is DAX in Power BI?", "hint": "Data Analysis Expressions, measures, calculated columns.",
         "a": "DAX (Data Analysis Expressions) is a formula language in Power BI for creating custom calculations. Measures are dynamic calculations evaluated at query time based on filter context (e.g., SUMX, CALCULATE, AVERAGEX). Calculated columns are computed row-by-row and stored in the model. Key functions: CALCULATE (modify filter context), FILTER, ALL (remove filters), RELATED (cross-table lookup), time intelligence (SAMEPERIODLASTYEAR). Measures are preferred over calculated columns for performance."},
        {"q": "How do you create relationships between tables in Power BI?", "hint": "Star schema, cardinality, cross-filter direction.",
         "a": "In Power BI, drag a column from one table to a matching column in another to create a relationship. Best practice: use a star schema with fact tables (transactions) connected to dimension tables (products, dates) via foreign keys. Cardinality types: one-to-many (most common), one-to-one, many-to-many (use cautiously). Cross-filter direction: single (dimension filters fact) or both (bidirectional, can cause ambiguity). Manage in Model view."},
    ],
}

# HR / Behavioral questions (universal, with answers)
HR_QUESTIONS = [
    {"q": "Tell me about yourself.", "hint": "2-minute pitch: background, skills, why this role.",
     "a": "Structure your answer as: Present (current role/education and key skills), Past (relevant experience and achievements), Future (why this role excites you). Example: 'I'm a Computer Science student passionate about AI and web development. I've built projects using Python, React, and machine learning, including a resume parser using RAG. I'm excited about this internship because it aligns with my goal to build AI-powered products that solve real problems.'"},
    {"q": "Why do you want this internship?", "hint": "Company research, role alignment, career goals.",
     "a": "Show genuine interest by mentioning: (1) What you know about the company's products/mission. (2) How the role matches your skills and interests. (3) What specific skills you want to develop. (4) How this fits your career goals. Avoid generic answers. Example: 'I admire how your company uses AI to solve X. My experience with Python and ML aligns well with this role, and I'm eager to learn production-level deployment practices from your team.'"},
    {"q": "What are your strengths and weaknesses?", "hint": "Be honest, show self-awareness, growth mindset.",
     "a": "For strengths, pick 2-3 relevant to the role with examples: 'I'm a fast learner — I taught myself React in 2 weeks and built a full-stack app.' For weaknesses, be honest but show improvement: 'I used to struggle with time estimation for projects, but I've started breaking tasks into smaller chunks and using project management tools, which has significantly improved my delivery timelines.' Never say 'I'm a perfectionist' — it sounds rehearsed."},
    {"q": "Describe a challenging project you worked on.", "hint": "STAR method: Situation, Task, Action, Result.",
     "a": "Use the STAR method: Situation (set the context), Task (your responsibility), Action (specific steps you took), Result (measurable outcome). Example: 'During my college project (S), I needed to build a resume matching system (T). I researched RAG pipelines, implemented semantic search using sentence-transformers, and built a FastAPI backend (A). The system achieved 85%+ accuracy in matching resumes to relevant job descriptions (R).'"},
    {"q": "Where do you see yourself in 5 years?", "hint": "Career growth, skill development, industry impact.",
     "a": "Show ambition while being realistic. Focus on skill growth and impact, not job titles. Example: 'In 5 years, I want to be a proficient full-stack developer with deep expertise in AI/ML. I'd love to lead projects that apply AI to real-world problems. I plan to contribute to open-source, stay updated with emerging tech, and possibly mentor junior developers. This internship is the first step toward building that expertise.'"},
    {"q": "How do you handle tight deadlines or pressure?", "hint": "Time management, prioritization, staying calm.",
     "a": "Show your approach: (1) Break the task into smaller, prioritized chunks. (2) Focus on the most critical features first (MVP approach). (3) Communicate early if the deadline is unrealistic. (4) Stay calm and avoid multitasking. Give an example: 'During my final semester, I had 3 project deadlines in one week. I created a priority matrix, focused on high-impact tasks first, and communicated with my team about task distribution. We delivered all projects on time.'"},
    {"q": "Describe a time you worked in a team.", "hint": "Collaboration, conflict resolution, communication.",
     "a": "Highlight collaboration and your role. Example: 'In a hackathon team of 4, we built a chatbot in 24 hours. I took ownership of the backend API while others handled frontend and ML. When we had a disagreement about the tech stack, I suggested we evaluate options against our time constraint, and we agreed on FastAPI for speed. We won second place. I learned that clear communication and defined responsibilities are key to team success.'"},
    {"q": "Why should we hire you over other candidates?", "hint": "Unique skills, enthusiasm, cultural fit.",
     "a": "Focus on what makes you unique without putting others down. Mention: (1) Your specific technical skills relevant to the role. (2) Projects that demonstrate your ability. (3) Your eagerness to learn and contribute. Example: 'I bring a combination of hands-on project experience in AI and web development, a strong foundation in Python and React, and genuine enthusiasm for this domain. My resume matching project shows I can build end-to-end AI applications independently.'"},
    {"q": "Do you have any questions for us?", "hint": "Always ask! Team culture, growth, tech stack, mentorship.",
     "a": "Always ask 2-3 thoughtful questions: (1) 'What does a typical day look like for interns on your team?' (2) 'What tech stack does the team use and are there opportunities to learn new technologies?' (3) 'How do you measure success for this internship role?' (4) 'What's the team culture like — do you do code reviews, pair programming?' Avoid asking about salary/benefits in the first interview. Show genuine curiosity about the role and company."},
]


@app.post("/internships/{internship_id}/interview-questions", tags=["Interview Prep"])
def generate_interview_questions(
    internship_id: str,
    db: Session = Depends(get_db),
    current_user: DBUser = Depends(get_current_user)
):
    """Generate personalised interview questions based on user's resume and selected internship."""
    job = next((j for j in INTERNSHIP_DATABASE if j["id"] == internship_id), None)
    if not job:
        raise HTTPException(status_code=404, detail="Internship not found")

    profile = db.query(DBUserProfile).filter(DBUserProfile.user_id == current_user.id).first()
    if not profile:
        raise HTTPException(status_code=400, detail="Please upload your resume first.")

    user_skills = set(s.lower() for s in (profile.skills or []))
    required_skills = job.get("required_skills", [])

    # 1. Technical questions — from matched + required skills
    technical = []
    for skill in required_skills:
        questions = SKILL_QUESTIONS.get(skill, [])
        is_matched = skill.lower() in user_skills
        for q_obj in questions:
            technical.append({
                "question": q_obj["q"],
                "hint": q_obj["hint"],
                "answer": q_obj["a"],
                "skill": skill,
                "matched": is_matched,
                "category": "technical"
            })

    # 2. HR questions
    hr = [{"question": q["q"], "hint": q["hint"], "answer": q["a"], "skill": None, "matched": False, "category": "hr"} for q in HR_QUESTIONS]

    # 3. Role-specific situational questions
    situational = [
        {"question": f"What interests you about the {job['title']} role at {job['company']}?",
         "hint": "Research the company, connect your skills to their mission.", "skill": None, "matched": False, "category": "situational"},
        {"question": f"How would your experience with {', '.join(required_skills[:3])} help you in this role?",
         "hint": "Give specific project examples demonstrating these skills.", "skill": None, "matched": False, "category": "situational"},
        {"question": f"This role involves {job.get('description', 'various tasks')}. Describe a similar project you've done.",
         "hint": "Use STAR method, quantify your impact.", "skill": None, "matched": False, "category": "situational"},
    ]

    # 4. Preparation roadmap
    matched_skills = [s for s in required_skills if s.lower() in user_skills]
    missing_skills = [s for s in required_skills if s.lower() not in user_skills]

    roadmap = {
        "strong_areas": matched_skills,
        "areas_to_improve": missing_skills,
        "preparation_tips": [
            f"Review fundamentals of {', '.join(required_skills[:3])}.",
            "Practice coding problems on LeetCode/HackerRank for 30 min daily.",
            f"Build a small project using {required_skills[0] if required_skills else 'relevant tech'}.",
            "Prepare 2-3 STAR stories from your projects/experience.",
            "Research the company's products, culture, and recent news.",
            "Practice mock interviews with a friend or in front of a mirror.",
        ],
        "recommended_topics": [
            f"{skill} — core concepts, best practices, common pitfalls" for skill in required_skills[:5]
        ],
    }

    return {
        "status": "success",
        "internship": {"id": job["id"], "title": job["title"], "company": job["company"]},
        "technical_questions": technical,
        "hr_questions": hr,
        "situational_questions": situational,
        "roadmap": roadmap,
        "summary": {
            "total_questions": len(technical) + len(hr) + len(situational),
            "technical_count": len(technical),
            "hr_count": len(hr),
            "situational_count": len(situational),
            "skills_matched": len(matched_skills),
            "skills_to_learn": len(missing_skills),
        }
    }

# ==============================================================================
# 4e. RESUME ANALYZER (ATS Score + Feedback)
# ==============================================================================

@app.post("/resume/analyze", tags=["Resume Analyzer"])
def analyze_resume(
    db: Session = Depends(get_db),
    current_user: DBUser = Depends(get_current_user)
):
    """Analyze resume quality — ATS score, section feedback, and improvement tips."""
    profile = db.query(DBUserProfile).filter(DBUserProfile.user_id == current_user.id).first()
    if not profile:
        raise HTTPException(status_code=400, detail="Upload your resume first from the Dashboard.")

    resume = db.query(DBResume).filter(DBResume.user_id == current_user.id).order_by(DBResume.id.desc()).first()
    # Build resume_text from extracted_data JSON or profile fields
    resume_text = ""
    if resume and resume.extracted_data:
        data = resume.extracted_data if isinstance(resume.extracted_data, dict) else {}
        resume_text = str(data)
    if not resume_text and profile:
        resume_text = " ".join(filter(None, [
            profile.summary or "", str(profile.skills or ""),
            profile.experience or "", profile.education or "",
            profile.projects or "", profile.certifications or ""
        ]))

    skills = profile.skills or []
    if isinstance(skills, str):
        skills = [s.strip() for s in skills.split(",") if s.strip()]

    # experience and education are Text fields, convert to list-like for counting
    exp_raw = profile.experience or ""
    edu_raw = profile.education or ""
    experience_entries = [e.strip() for e in exp_raw.split("\n") if e.strip()] if exp_raw else []
    education_entries = [e.strip() for e in edu_raw.split("\n") if e.strip()] if edu_raw else []
    name = current_user.full_name or ""
    email = profile.resume_email or current_user.email or ""

    # ── Section scores ──
    sections = {}
    total_score = 0

    # 1. Contact Info (15 pts)
    contact_score = 0
    contact_tips = []
    if name and len(name) > 2: contact_score += 5
    else: contact_tips.append("Add your full name at the top of your resume.")
    if email and "@" in email: contact_score += 5
    else: contact_tips.append("Include a professional email address.")
    if resume_text and any(w in resume_text.lower() for w in ["linkedin", "github", "portfolio", "phone", "+91", "+1"]):
        contact_score += 5
    else: contact_tips.append("Add LinkedIn/GitHub profile links and phone number.")
    sections["contact_info"] = {"score": contact_score, "max": 15, "tips": contact_tips, "label": "Contact Information"}
    total_score += contact_score

    # 2. Skills (20 pts)
    skill_score = 0
    skill_tips = []
    if len(skills) >= 8: skill_score = 20
    elif len(skills) >= 5: skill_score = 15
    elif len(skills) >= 3: skill_score = 10
    elif len(skills) >= 1: skill_score = 5
    else: skill_tips.append("Add at least 5-8 relevant technical skills.")
    if len(skills) < 5:
        skill_tips.append(f"You have {len(skills)} skills listed. Aim for 8-12 for better ATS matching.")
    # Check for mix of hard and soft skills
    soft_skills_found = any(s.lower() in ["communication","teamwork","leadership","problem solving","time management"] for s in skills)
    if not soft_skills_found:
        skill_tips.append("Consider adding 2-3 soft skills (Communication, Leadership, Problem Solving).")
    sections["skills"] = {"score": skill_score, "max": 20, "tips": skill_tips, "label": "Skills"}
    total_score += skill_score

    # 3. Experience (25 pts)
    exp_score = 0
    exp_tips = []
    if len(experience_entries) >= 3: exp_score = 20
    elif len(experience_entries) >= 2: exp_score = 15
    elif len(experience_entries) >= 1: exp_score = 10
    else:
        exp_tips.append("Add work experience, internships, or project experience.")
        exp_score = 0
    # Check for action verbs and metrics
    action_verbs = ["developed","built","designed","implemented","managed","led","created","improved","increased","reduced","launched","deployed","optimized","automated"]
    exp_text = exp_raw.lower()
    verbs_found = sum(1 for v in action_verbs if v in exp_text)
    if verbs_found >= 3: exp_score += 5
    elif verbs_found >= 1: exp_score += 3
    else: exp_tips.append("Use strong action verbs: Developed, Built, Implemented, Improved, Led.")
    if not any(c.isdigit() for c in exp_text):
        exp_tips.append("Quantify achievements with numbers: 'Improved performance by 40%', 'Managed team of 5'.")
    sections["experience"] = {"score": min(exp_score, 25), "max": 25, "tips": exp_tips, "label": "Experience"}
    total_score += min(exp_score, 25)

    # 4. Education (15 pts)
    edu_score = 0
    edu_tips = []
    if len(education_entries) >= 1: edu_score = 10
    else: edu_tips.append("Add your education details (degree, university, year).")
    edu_text = edu_raw.lower()
    if any(w in edu_text for w in ["gpa","cgpa","percentage","grade","honors","distinction"]):
        edu_score += 5
    else: edu_tips.append("Include your GPA/CGPA or percentage if it's strong (above 7.0/10 or 70%).")
    sections["education"] = {"score": edu_score, "max": 15, "tips": edu_tips, "label": "Education"}
    total_score += edu_score

    # 5. Formatting & Length (15 pts)
    fmt_score = 0
    fmt_tips = []
    text_len = len(resume_text)
    if 500 <= text_len <= 5000:
        fmt_score += 8
    elif text_len > 5000:
        fmt_score += 4
        fmt_tips.append("Resume is very long. Keep it to 1-2 pages for internships.")
    elif text_len > 200:
        fmt_score += 4
        fmt_tips.append("Resume seems short. Add more detail to experience and projects.")
    else:
        fmt_tips.append("Resume has very little content. Expand all sections.")
    # Check for section headers
    header_words = ["experience","education","skills","projects","certifications","objective","summary","achievements"]
    headers_found = sum(1 for h in header_words if h in resume_text.lower())
    if headers_found >= 4: fmt_score += 7
    elif headers_found >= 2: fmt_score += 4
    else: fmt_tips.append("Use clear section headers: Experience, Education, Skills, Projects.")
    sections["formatting"] = {"score": fmt_score, "max": 15, "tips": fmt_tips, "label": "Formatting & Structure"}
    total_score += fmt_score

    # 6. ATS Keywords (10 pts)
    ats_score = 0
    ats_tips = []
    ats_keywords = ["python","javascript","react","sql","api","database","machine learning","data","cloud",
                    "git","docker","agile","testing","deployment","ci/cd","rest","framework"]
    keywords_found = [k for k in ats_keywords if k in resume_text.lower()]
    if len(keywords_found) >= 6: ats_score = 10
    elif len(keywords_found) >= 4: ats_score = 7
    elif len(keywords_found) >= 2: ats_score = 4
    else: ats_score = 1
    if len(keywords_found) < 6:
        missing_kw = [k for k in ats_keywords if k not in resume_text.lower()][:5]
        ats_tips.append(f"Consider adding trending keywords: {', '.join(missing_kw)}.")
    sections["ats_keywords"] = {"score": ats_score, "max": 10, "tips": ats_tips, "label": "ATS Keywords"}
    total_score += ats_score

    # Overall grade
    if total_score >= 85: grade = "A+"
    elif total_score >= 75: grade = "A"
    elif total_score >= 65: grade = "B+"
    elif total_score >= 55: grade = "B"
    elif total_score >= 45: grade = "C"
    else: grade = "D"

    # Top improvement tips
    all_tips = []
    for sec in sections.values():
        all_tips.extend(sec["tips"])

    return {
        "status": "success",
        "overall_score": total_score,
        "max_score": 100,
        "grade": grade,
        "sections": sections,
        "top_tips": all_tips[:8],
        "profile_summary": {
            "name": name,
            "email": email,
            "skills_count": len(skills),
            "experience_count": len(experience_entries),
            "education_count": len(education_entries),
            "resume_length": text_len,
        }
    }

# ==============================================================================
# 4f. MOCK INTERVIEW
# ==============================================================================

@app.post("/mock-interview/start", tags=["Mock Interview"])
def start_mock_interview(
    internship_id: str = None,
    count: int = 10,
    db: Session = Depends(get_db),
    current_user: DBUser = Depends(get_current_user)
):
    """Start a mock interview session — returns randomized questions with timer settings."""
    import random

    profile = db.query(DBUserProfile).filter(DBUserProfile.user_id == current_user.id).first()

    all_questions = []

    # If internship selected, use skill-matched questions
    if internship_id:
        job = next((j for j in INTERNSHIP_DATABASE if j["id"] == internship_id), None)
        if job:
            for skill in job.get("required_skills", []):
                for q_obj in SKILL_QUESTIONS.get(skill, []):
                    all_questions.append({
                        "question": q_obj["q"],
                        "answer": q_obj["a"],
                        "hint": q_obj["hint"],
                        "category": "technical",
                        "skill": skill,
                        "time_limit": 120,
                    })

    # If not enough questions, add from all skills
    if len(all_questions) < count:
        for skill, qs in SKILL_QUESTIONS.items():
            for q_obj in qs:
                q_entry = {
                    "question": q_obj["q"],
                    "answer": q_obj["a"],
                    "hint": q_obj["hint"],
                    "category": "technical",
                    "skill": skill,
                    "time_limit": 120,
                }
                if q_entry not in all_questions:
                    all_questions.append(q_entry)

    # Add HR questions
    for q_obj in HR_QUESTIONS:
        all_questions.append({
            "question": q_obj["q"],
            "answer": q_obj["a"],
            "hint": q_obj["hint"],
            "category": "hr",
            "skill": None,
            "time_limit": 90,
        })

    # Shuffle and pick `count` questions
    random.shuffle(all_questions)
    selected = all_questions[:min(count, len(all_questions))]

    # Add question numbers
    for i, q in enumerate(selected):
        q["id"] = i + 1

    return {
        "status": "success",
        "session": {
            "total_questions": len(selected),
            "estimated_time_minutes": round(sum(q["time_limit"] for q in selected) / 60, 1),
            "questions": selected,
        }
    }

# ==============================================================================
# 4g. AI INTERVIEW PREP CHATBOT
# ==============================================================================

CHATBOT_INTENTS = {
    "greet": {
        "keywords": ["hello","hi","hey","good morning","good afternoon","good evening","namaste"],
        "response": "Hello {name}! 👋 I'm your AI Interview Preparation Assistant. I can help you with:\n\n"
                    "• 🎯 Interview tips for your skills\n"
                    "• 💡 How to answer common questions\n"
                    "• 📝 Resume improvement advice\n"
                    "• 🛠️ Technical concept explanations\n"
                    "• 🤝 HR & behavioral question guidance\n\n"
                    "What would you like help with?"
    },
    "tell_me_about_yourself": {
        "keywords": ["tell me about yourself","introduce yourself","about me","self introduction","how to introduce"],
        "response": "Here's a structured way to answer 'Tell me about yourself', {name}:\n\n"
                    "**1. Present** — Start with your current role/education:\n"
                    "\"I'm {name}, {summary_line}\"\n\n"
                    "**2. Past** — Highlight relevant experience:\n"
                    "{experience_tip}\n\n"
                    "**3. Future** — Show enthusiasm:\n"
                    "\"I'm excited to apply my skills in {top_skills} to contribute to your team.\"\n\n"
                    "**💡 Tips:**\n"
                    "• Keep it under 2 minutes\n"
                    "• Focus on what's relevant to the role\n"
                    "• End with why you're interested in THIS opportunity"
    },
    "strengths": {
        "keywords": ["strength","strengths","strong point","what are you good at","best quality"],
        "response": "Based on your profile, {name}, here are strengths you can highlight:\n\n"
                    "**Technical Strengths:**\n{skills_list}\n\n"
                    "**How to answer:**\n"
                    "1. Pick 2-3 strengths relevant to the job\n"
                    "2. Back each with a specific example\n"
                    "3. Use the STAR method (Situation, Task, Action, Result)\n\n"
                    "**Example:** \"One of my strengths is {top_skill}. For instance, {star_example}\""
    },
    "weaknesses": {
        "keywords": ["weakness","weaknesses","weak point","area of improvement","shortcoming"],
        "response": "Great question, {name}! Here's how to handle 'What's your weakness?':\n\n"
                    "**Formula:** Real weakness + Steps you're taking to improve\n\n"
                    "**Good examples:**\n"
                    "• \"I sometimes overthink solutions. I've started using timeboxing to decide faster.\"\n"
                    "• \"Public speaking used to be challenging. I've been presenting in team meetings to improve.\"\n"
                    "• \"I can be too detail-oriented. I now set priorities to balance quality with deadlines.\"\n\n"
                    "**❌ Avoid:**\n"
                    "• \"I'm a perfectionist\" (overused)\n"
                    "• \"I work too hard\" (not genuine)\n"
                    "• Anything that's a core job requirement"
    },
    "why_this_company": {
        "keywords": ["why this company","why do you want to work","why join","why should we hire","why us"],
        "response": "Here's a framework for 'Why this company?', {name}:\n\n"
                    "**Structure:**\n"
                    "1. **Company Research** — Mention something specific (product, mission, culture)\n"
                    "2. **Skill Alignment** — Connect your skills ({top_skills}) to their needs\n"
                    "3. **Growth** — Show how you'll grow there AND contribute\n\n"
                    "**Template:**\n"
                    "\"I admire [company]'s work in [specific area]. With my skills in {top_skills}, "
                    "I can contribute to [specific project/team]. I'm also excited about the opportunity "
                    "to grow in [area] alongside your team.\"\n\n"
                    "**💡 Always research the company before the interview!**"
    },
    "salary": {
        "keywords": ["salary","compensation","pay","package","ctc","expected salary","salary negotiation"],
        "response": "Here's how to handle salary questions, {name}:\n\n"
                    "**If asked early:**\n"
                    "\"I'd like to learn more about the role first. I'm open to discussing a fair package based on the responsibilities.\"\n\n"
                    "**If pressed:**\n"
                    "\"Based on my research and skills in {top_skills}, I'd expect a range of [X-Y]. But I'm flexible and more focused on growth opportunities.\"\n\n"
                    "**Tips:**\n"
                    "• Research market rates on Glassdoor/Levels.fyi\n"
                    "• Give a range, not a fixed number\n"
                    "• For internships, ask about stipend + learning opportunities\n"
                    "• Never share previous salary"
    },
    "star_method": {
        "keywords": ["star method","star technique","behavioral question","situation task action result","how to answer behavioral"],
        "response": "The STAR Method is essential for behavioral questions, {name}!\n\n"
                    "**S — Situation:** Set the scene (where, when, context)\n"
                    "**T — Task:** What was your responsibility?\n"
                    "**A — Action:** What specifically did YOU do?\n"
                    "**R — Result:** What was the outcome? (use numbers!)\n\n"
                    "**Example with your skills:**\n"
                    "\"While working on a project using {top_skill}, (S) our team faced a tight deadline. "
                    "(T) I was responsible for the core implementation. "
                    "(A) I broke the task into sprints and used {second_skill} to optimize the workflow. "
                    "(R) We delivered 2 days early with 95% test coverage.\"\n\n"
                    "**Use STAR for:** Leadership, teamwork, conflict, failure, achievement questions"
    },
    "technical_tips": {
        "keywords": ["technical interview","coding interview","technical round","dsa","data structure","algorithm","coding tips"],
        "response": "Technical interview tips based on your skills ({top_skills}), {name}:\n\n"
                    "**Before the interview:**\n"
                    "• Review fundamentals of {top_skill}\n"
                    "• Practice 2-3 problems daily on LeetCode/HackerRank\n"
                    "• Revise common patterns: arrays, strings, trees, DP\n\n"
                    "**During the interview:**\n"
                    "1. **Clarify** the problem — ask questions\n"
                    "2. **Think aloud** — explain your approach\n"
                    "3. **Start simple** — brute force first, then optimize\n"
                    "4. **Test** your code with examples\n"
                    "5. **Discuss** time & space complexity\n\n"
                    "**💡 It's okay to say:** \"Let me think about this for a moment\""
    },
    "resume_tips": {
        "keywords": ["resume tip","improve resume","resume advice","cv tip","resume help","ats"],
        "response": "Resume improvement tips for you, {name}:\n\n"
                    "**Your current profile:**\n"
                    "• Skills: {skills_count} listed\n"
                    "• Format: Use our 📊 Resume Analyzer for a detailed score!\n\n"
                    "**Quick tips:**\n"
                    "1. Use **action verbs**: Built, Developed, Improved, Led\n"
                    "2. **Quantify** results: 'Improved speed by 40%'\n"
                    "3. Keep it to **1 page** for internships\n"
                    "4. Tailor skills to each job description\n"
                    "5. Add **projects** with GitHub links\n"
                    "6. Include relevant **keywords** for ATS\n\n"
                    "**Go to Resume Score (📊) in the navbar for a full analysis!**"
    },
    "project_discussion": {
        "keywords": ["project","tell me about your project","discuss project","explain project","project question"],
        "response": "Here's how to discuss your projects in interviews, {name}:\n\n"
                    "**Structure (2-3 minutes per project):**\n"
                    "1. **What** — One-line description of the project\n"
                    "2. **Why** — The problem it solved\n"
                    "3. **How** — Tech stack & your role (mention {top_skills})\n"
                    "4. **Challenges** — What was difficult & how you solved it\n"
                    "5. **Result** — Impact, metrics, or learnings\n\n"
                    "**Tips:**\n"
                    "• Prepare 2-3 projects in detail\n"
                    "• Be ready for deep-dive follow-up questions\n"
                    "• If it's a team project, clarify YOUR contribution\n"
                    "• Mention testing/deployment if applicable"
    },
    "confidence": {
        "keywords": ["nervous","anxiety","confident","confidence","scared","fear","stage fright","calm"],
        "response": "Interview confidence tips for you, {name}! 💪\n\n"
                    "**Before:**\n"
                    "• Practice with our ⏱️ Mock Interview feature\n"
                    "• Prepare answers for top 10 common questions\n"
                    "• Research the company thoroughly\n"
                    "• Get 7-8 hours of sleep\n\n"
                    "**During:**\n"
                    "• Take a deep breath before answering\n"
                    "• It's okay to pause and think\n"
                    "• Maintain eye contact (or camera for video)\n"
                    "• Sit upright — posture affects confidence\n"
                    "• Smile naturally\n\n"
                    "**Remember:** The interviewer WANTS you to succeed! They're looking for a teammate, not trying to catch you out."
    },
    "thank_you": {
        "keywords": ["thank","thanks","thank you","thx","appreciate"],
        "response": "You're welcome, {name}! 😊 Remember:\n\n"
                    "• Practice makes perfect — use our ⏱️ Mock Interview\n"
                    "• Check your 📊 Resume Score for improvements\n"
                    "• Review 🎓 Interview Prep questions\n\n"
                    "Good luck with your interviews! You've got this! 🚀"
    },
}

FALLBACK_RESPONSES = [
    "That's an interesting question, {name}! While I specialize in interview prep, here are some general tips:\n\n"
    "• Research the company and role thoroughly\n"
    "• Practice with our Mock Interview feature ⏱️\n"
    "• Use the STAR method for behavioral questions\n"
    "• Review your skills: {top_skills}\n\n"
    "Try asking about: 'How to introduce myself', 'STAR method', 'technical interview tips', or 'resume advice'",

    "I'm not sure about that specific topic, {name}, but here's what I can help with:\n\n"
    "• 🎯 'Tell me about yourself' preparation\n"
    "• 💪 Strengths & weaknesses answers\n"
    "• 🏢 'Why this company?' framework\n"
    "• 📝 Resume improvement tips\n"
    "• 🛠️ Technical interview strategies\n"
    "• 🤝 HR question guidance\n\n"
    "What would you like to know?",
]


class ChatRequest(BaseModel):
    message: str


@app.post("/chatbot/message", tags=["Chatbot"])
def chatbot_message(
    req: ChatRequest,
    db: Session = Depends(get_db),
    current_user: DBUser = Depends(get_current_user)
):
    """AI Interview Prep Chatbot — uses resume context for personalized responses."""
    import random

    # Load user context
    profile = db.query(DBUserProfile).filter(DBUserProfile.user_id == current_user.id).first()
    name = current_user.full_name.split()[0] if current_user.full_name else "there"
    skills = []
    skills_count = 0
    summary_line = "a passionate learner"
    experience_tip = "Highlight any internships, projects, or coursework."
    top_skill = "problem solving"
    second_skill = "teamwork"
    top_skills = "your core skills"

    if profile:
        skills = profile.skills or []
        if isinstance(skills, str):
            skills = [s.strip() for s in skills.split(",") if s.strip()]
        skills_count = len(skills)
        top_skills = ", ".join(skills[:4]) if skills else "your core skills"
        top_skill = skills[0] if skills else "problem solving"
        second_skill = skills[1] if len(skills) > 1 else "collaboration"
        if profile.summary:
            summary_line = profile.summary[:100]
        if profile.experience:
            exp_lines = [e.strip() for e in profile.experience.split("\n") if e.strip()]
            if exp_lines:
                experience_tip = f"Mention: \"{exp_lines[0][:80]}...\""

    skills_list = "\n".join(f"• {s}" for s in skills[:6]) if skills else "• Upload your resume to see personalized suggestions!"
    star_example = f"I used {top_skill} to solve a key challenge, resulting in improved efficiency."

    # Context dict for formatting
    ctx = {
        "name": name, "top_skills": top_skills, "top_skill": top_skill,
        "second_skill": second_skill, "skills_list": skills_list,
        "skills_count": skills_count, "summary_line": summary_line,
        "experience_tip": experience_tip, "star_example": star_example,
    }

    user_msg = req.message.lower().strip()

    # 1. Intent matching
    best_intent = None
    best_score = 0
    for intent_name, intent_data in CHATBOT_INTENTS.items():
        for keyword in intent_data["keywords"]:
            if keyword in user_msg:
                score = len(keyword)  # longer keyword = more specific match
                if score > best_score:
                    best_score = score
                    best_intent = intent_name

    if best_intent:
        response = CHATBOT_INTENTS[best_intent]["response"].format(**ctx)
        intent = best_intent
    else:
        # 2. Skill-specific question — check if user asks about a skill they have
        asked_skill = None
        for skill in skills:
            if skill.lower() in user_msg:
                asked_skill = skill
                break

        if asked_skill:
            # Pull questions from SKILL_QUESTIONS if available
            skill_qs = SKILL_QUESTIONS.get(asked_skill, [])
            if skill_qs:
                picked = random.sample(skill_qs, min(3, len(skill_qs)))
                q_text = "\n".join(f"**Q{i+1}:** {q['q']}\n💡 Hint: {q['hint']}" for i, q in enumerate(picked))
                response = (f"Here are some {asked_skill} interview questions for you, {name}:\n\n"
                           f"{q_text}\n\n"
                           f"**Go to 🎓 Interview Prep for more {asked_skill} questions with full answers!**")
            else:
                response = (f"Great that you know {asked_skill}, {name}! To prepare:\n\n"
                           f"• Review core concepts and common patterns\n"
                           f"• Practice coding problems related to {asked_skill}\n"
                           f"• Be ready to explain a project where you used {asked_skill}\n"
                           f"• Prepare for 'What did you build with {asked_skill}?' questions\n\n"
                           f"Check 🎓 Interview Prep for targeted questions!")
            intent = f"skill_{asked_skill}"
        else:
            # 3. Semantic search over intents using model
            try:
                model = vector_store.model
                intent_texts = []
                intent_keys = []
                for k, v in CHATBOT_INTENTS.items():
                    intent_texts.extend(v["keywords"])
                    intent_keys.extend([k] * len(v["keywords"]))

                q_emb = model.encode([user_msg])
                i_embs = model.encode(intent_texts)
                sims = cosine_similarity(q_emb, i_embs)[0]
                best_idx = int(np.argmax(sims))

                if sims[best_idx] > 0.35:
                    matched_intent = intent_keys[best_idx]
                    response = CHATBOT_INTENTS[matched_intent]["response"].format(**ctx)
                    intent = matched_intent
                else:
                    response = random.choice(FALLBACK_RESPONSES).format(**ctx)
                    intent = "fallback"
            except Exception:
                response = random.choice(FALLBACK_RESPONSES).format(**ctx)
                intent = "fallback"

    return {
        "status": "success",
        "response": response,
        "intent": intent,
        "user_context": {
            "name": name,
            "skills_count": skills_count,
            "has_profile": profile is not None,
        }
    }


@app.get("/chatbot/suggestions", tags=["Chatbot"])
def chatbot_suggestions(
    db: Session = Depends(get_db),
    current_user: DBUser = Depends(get_current_user)
):
    """Get personalized quick-reply suggestions for the chatbot."""
    profile = db.query(DBUserProfile).filter(DBUserProfile.user_id == current_user.id).first()
    name = current_user.full_name.split()[0] if current_user.full_name else "there"

    suggestions = [
        "How do I introduce myself?",
        "What are good strengths to mention?",
        "How to handle weakness questions?",
        "Explain the STAR method",
        "Technical interview tips",
        "Resume improvement advice",
    ]

    if profile and profile.skills:
        skills = profile.skills if isinstance(profile.skills, list) else []
        if skills:
            suggestions.insert(2, f"Interview questions about {skills[0]}")
            if len(skills) > 2:
                suggestions.insert(3, f"How to discuss {skills[1]} in interviews?")

    return {
        "status": "success",
        "greeting": f"Hi {name}! 👋 How can I help you prepare for your interview?",
        "suggestions": suggestions[:8],
    }

# ==============================================================================
# 5. RAG & PARSING UTILS
# ==============================================================================
class InternshipVectorStore:
    def __init__(self):
        self.model = SentenceTransformer("all-MiniLM-L6-v2")
        self.documents = [f"{j['title']} {' '.join(j['required_skills'])}" for j in INTERNSHIP_DATABASE]
        self.embeddings = self.model.encode(self.documents)

    def search(self, query_text: str, top_k: int = 3):
        query_vector = self.model.encode([query_text])
        similarities = cosine_similarity(query_vector, self.embeddings)[0]
        top_indices = np.argsort(similarities)[::-1][:top_k]
        return [{"job": INTERNSHIP_DATABASE[idx], "similarity_score": round(float(similarities[idx]), 4)} for idx in top_indices]

vector_store = InternshipVectorStore()

def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    if filename.endswith(".pdf") and fitz:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        return "\n".join([page.get_text() for page in doc])
    elif filename.endswith(".docx") and docx:
        with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
            tmp.write(file_bytes)
            tmp_path = tmp.name
        doc = docx.Document(tmp_path)
        os.unlink(tmp_path)
        return "\n".join([para.text for para in doc.paragraphs])
    return file_bytes.decode("utf-8", errors="ignore")

class ComprehensiveResumeParser:
    @staticmethod
    def extract_section(raw_text: str, target_headers: List[str]) -> Optional[str]:
        lines = raw_text.split('\n')
        section_blocks = []
        in_section = False
        major_headers = ["EDUCATION", "PROJECTS", "TECHNICAL SKILLS", "CERTIFICATIONS",
                         "LANGUAGES", "WORK EXPERIENCE", "EXPERIENCE", "SUMMARY",
                         "PROFESSIONAL SUMMARY", "SKILLS", "INTERNSHIP", "INTERNSHIPS"]
        for line in lines:
            clean_line = line.strip()
            if not clean_line:
                continue
            if clean_line.upper() in target_headers:
                in_section = True
                section_blocks.append(clean_line)
                continue
            if in_section and clean_line.upper() in major_headers and clean_line.upper() not in target_headers:
                break
            if in_section:
                section_blocks.append(clean_line)
        return "\n".join(section_blocks[1:]).strip() if section_blocks else None

    @staticmethod
    def parse_full_resume(raw_text: str) -> Dict[str, Any]:
        email_match = re.search(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+", raw_text)
        # Clean simple phone regex — works for Indian & international numbers
        phone_match = re.search(r"[\+]?[(]?[0-9]{1,4}[)]?[-\s\.]?[(]?[0-9]{2,4}[)]?[-\s\.]?[0-9]{3,4}[-\s\.]?[0-9]{3,4}", raw_text)
        master_skills_list = [
            # AI / ML
            "Python", "PyTorch", "TensorFlow", "Hugging Face", "LLMs", "RAG", "LangChain",
            "Scikit-Learn", "Keras", "OpenCV", "NLTK", "Transformers",
            # Backend / APIs
            "FastAPI", "Django", "Flask", "REST APIs", "GraphQL", "Node.js", "Express.js",
            "PostgreSQL", "MySQL", "MongoDB", "Redis", "SQLite",
            # DevOps / Cloud
            "Docker", "AWS", "Kubernetes", "GCP", "Azure", "Linux",
            # Frontend / Web
            "React.js", "JavaScript", "TypeScript", "HTML", "CSS", "Vue.js", "Next.js",
            "Tailwind", "Bootstrap", "Redux",
            # Data
            "Pandas", "SQL", "NumPy", "Data Visualization", "Streamlit", "Power BI",
            "Tableau", "Excel", "Spark", "Hadoop",
            # Other languages
            "Java", "PHP", "Go", "Rust", "Kotlin", "Swift",
        ]
        def safe_skill_search(skill, text):
            try:
                return bool(re.search(re.escape(skill), text, re.IGNORECASE))
            except re.error:
                return False
        found_skills = [s for s in master_skills_list if safe_skill_search(s, raw_text)]
        return {
            "contact_info": {
                "email": email_match.group(0) if email_match else None,
                "phone_number": phone_match.group(0) if phone_match else None
            },
            "summary":        ComprehensiveResumeParser.extract_section(raw_text, ["SUMMARY", "PROFESSIONAL SUMMARY"]),
            "education":      ComprehensiveResumeParser.extract_section(raw_text, ["EDUCATION"]),
            "internships":    ComprehensiveResumeParser.extract_section(raw_text, ["INTERNSHIP", "INTERNSHIPS", "WORK EXPERIENCE", "EXPERIENCE"]),
            "projects":       ComprehensiveResumeParser.extract_section(raw_text, ["PROJECTS"]),
            "certifications": ComprehensiveResumeParser.extract_section(raw_text, ["CERTIFICATIONS"]),
            "extracted_skills": found_skills
        }

# ==============================================================================
# 6. COVER LETTER GENERATOR
# ==============================================================================
def generate_cover_letter(user: DBUser, profile: DBUserProfile, job: dict) -> str:
    today = datetime.now().strftime("%d %B %Y")
    skills_str = ", ".join((profile.skills or [])[:6]) or "a diverse technical skill set"
    top_skills = ", ".join((profile.skills or [])[:3]) or "Python, data analysis, and software development"

    # Pick 1-2 lines from experience if available
    exp_snippet = ""
    if profile.experience:
        lines = [l.strip() for l in profile.experience.split("\n") if l.strip()]
        exp_snippet = lines[0] if lines else ""

    cover = f"""{today}

Hiring Manager
{job['company']}

Dear Hiring Manager,

I am writing to express my sincere interest in the {job['title']} position at {job['company']}. \
As a motivated and technically skilled candidate with hands-on experience in {top_skills}, \
I am confident that I can make a meaningful contribution to your team.

Throughout my academic and professional journey, I have developed strong proficiency in {skills_str}. \
{"Notably, " + exp_snippet + "." if exp_snippet else "I have applied these skills across multiple real-world projects and internships."} \
This background aligns closely with the requirements for the {job['title']} role, \
particularly the emphasis on {", ".join(job.get("required_skills", [])[:4])}.

I am especially drawn to {job['company']} because of its commitment to innovation and excellence. \
The {job['work_mode']} work model and {job['duration']} duration make this an ideal opportunity \
for me to contribute meaningfully while growing both technically and professionally.

{"My educational background, summarised as: " + profile.education.split(chr(10))[0] + ", has equipped me with a strong theoretical foundation to complement my practical skills." if profile.education else ""}

I am enthusiastic about the possibility of joining {job['company']} and would welcome the opportunity \
to discuss how my background, skills, and passion align with your team's goals.

Thank you sincerely for considering my application. I look forward to hearing from you.

Warm regards,
{user.full_name}
{user.email}
{profile.phone or ""}
"""
    return cover.strip()




# ==============================================================================
# 8. RESUME PARSING API (now also auto-fills user profile)
# ==============================================================================
@app.post("/resume/parse", tags=["Resume Parsing"])
async def upload_and_parse_resume(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: DBUser = Depends(get_current_user)
):
    try:
        # 1. Read & extract text
        file_bytes = await file.read()
        raw_text   = extract_text_from_file(file_bytes, file.filename)

        # 2. Parse
        parsed_data = ComprehensiveResumeParser.parse_full_resume(raw_text)

        # 3. RAG matching
        rag_query = f"Skills: {', '.join(parsed_data['extracted_skills'])} Context: {raw_text[:500]}"
        retrieved_matches = vector_store.search(rag_query, top_k=5)

        evaluated_matches = []
        candidate_skills_lower = set(s.lower() for s in parsed_data["extracted_skills"])
        for match in retrieved_matches:
            job = match["job"]
            req_skills = job.get("required_skills", [])
            matching_req = [s for s in req_skills if s.lower() in candidate_skills_lower]
            evaluated_matches.append({
                "job_id":          job["id"],
                "title":           job["title"],
                "company":         job["company"],
                "work_mode":       job.get("work_mode", ""),
                "duration":        job.get("duration", ""),
                "location":        job.get("location", ""),
                "description":     job.get("description", ""),
                "matched_skills":  matching_req,
                "composite_score": match["similarity_score"]
            })

        # 4. AUTO-FILL user profile from parsed resume data
        existing_profile = db.query(DBUserProfile).filter(DBUserProfile.user_id == current_user.id).first()
        profile_data = {
            "phone":          parsed_data["contact_info"].get("phone_number"),
            "summary":        parsed_data.get("summary"),
            "skills":         parsed_data.get("extracted_skills", []),
            "education":      parsed_data.get("education"),
            "experience":     parsed_data.get("internships"),
            "projects":       parsed_data.get("projects"),
            "certifications": parsed_data.get("certifications"),
            "resume_email":   parsed_data["contact_info"].get("email"),
        }
        if existing_profile:
            for key, val in profile_data.items():
                setattr(existing_profile, key, val)
        else:
            db.add(DBUserProfile(user_id=current_user.id, **profile_data))
        db.commit()

        # 5. Save full resume record
        full_response_data = {"parsed_resume": parsed_data, "rag_matches": evaluated_matches}
        db_resume = DBResume(user_id=current_user.id, filename=file.filename, extracted_data=full_response_data)
        db.add(db_resume)
        db.commit()
        db.refresh(db_resume)

        return {
            "status":             "success",
            "message":            "Successfully parsed, matched, profile updated, and saved to database.",
            "database_resume_id": db_resume.id,
            "data":               full_response_data
        }

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

# ==============================================================================
# 9. DOCUMENT Q&A APIs
# ==============================================================================

def chunk_text(text: str, chunk_size: int = 200, overlap: int = 60) -> List[str]:
    """Split text into smaller overlapping chunks for better retrieval accuracy."""
    # Clean up whitespace
    text = re.sub(r'\n{3,}', '\n\n', text.strip())
    text = re.sub(r' {2,}', ' ', text)

    sentences = re.split(r'(?<=[.!?])\s+|\n{2,}', text)
    sentences = [s.strip() for s in sentences if s.strip()]

    chunks, current = [], ""
    for sent in sentences:
        if len(current) + len(sent) > chunk_size and current:
            chunks.append(current.strip())
            # Keep last few words as overlap for context continuity
            words = current.split()
            overlap_words = max(3, overlap // 6)
            current = " ".join(words[-overlap_words:]) + " " + sent if len(words) > overlap_words else sent
        else:
            current = (current + " " + sent).strip()
    if current.strip():
        chunks.append(current.strip())
    return [c for c in chunks if len(c) > 15]


@app.post("/documents/upload", tags=["Document Q&A"])
async def upload_document(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: DBUser = Depends(get_current_user)
):
    """Upload a PDF/DOCX document, extract text, and chunk it for Q&A."""
    allowed = (".pdf", ".docx", ".txt")
    if not any(file.filename.lower().endswith(ext) for ext in allowed):
        raise HTTPException(status_code=400, detail="Only PDF, DOCX, and TXT files are supported.")

    try:
        file_bytes = await file.read()
        raw_text = extract_text_from_file(file_bytes, file.filename)

        if not raw_text or len(raw_text.strip()) < 20:
            raise HTTPException(status_code=400, detail="Could not extract meaningful text from this document.")

        chunks = chunk_text(raw_text)

        doc = DBDocument(
            user_id=current_user.id,
            filename=file.filename,
            extracted_text=raw_text,
            chunks=chunks,
            uploaded_at=datetime.utcnow().isoformat()
        )
        db.add(doc)
        db.commit()
        db.refresh(doc)

        return {
            "status": "success",
            "message": f"Document '{file.filename}' uploaded and processed.",
            "document_id": doc.id,
            "filename": doc.filename,
            "text_length": len(raw_text),
            "chunk_count": len(chunks),
            "preview": raw_text[:500] + ("..." if len(raw_text) > 500 else "")
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/documents", tags=["Document Q&A"])
def list_documents(
    db: Session = Depends(get_db),
    current_user: DBUser = Depends(get_current_user)
):
    """List all documents uploaded by the current user."""
    docs = db.query(DBDocument).filter(DBDocument.user_id == current_user.id).order_by(DBDocument.id.desc()).all()
    return {
        "documents": [
            {
                "id": d.id,
                "filename": d.filename,
                "text_length": len(d.extracted_text),
                "chunk_count": len(d.chunks) if d.chunks else 0,
                "uploaded_at": d.uploaded_at,
                "preview": d.extracted_text[:200] + ("..." if len(d.extracted_text) > 200 else "")
            }
            for d in docs
        ],
        "total": len(docs)
    }


class QuestionRequest(BaseModel):
    question: str


def _keyword_boost(question: str, chunk: str) -> float:
    """Calculate keyword overlap score between question and chunk."""
    stop_words = {"the","a","an","is","are","was","were","be","been","being","have","has","had",
                  "do","does","did","will","would","could","should","may","might","shall","can",
                  "in","on","at","to","for","of","with","by","from","about","between","through",
                  "what","how","why","when","where","who","which","this","that","these","those",
                  "and","or","but","not","no","it","its","i","me","my","we","our","you","your","they","their"}
    q_words = set(w.lower().strip("?.,!") for w in question.split() if len(w) > 2) - stop_words
    c_words_lower = chunk.lower()
    if not q_words:
        return 0.0
    matches = sum(1 for w in q_words if w in c_words_lower)
    return matches / len(q_words)


def _extract_best_sentences(question: str, chunk_text: str, model, max_sentences: int = 3) -> str:
    """Extract the most relevant sentences from a chunk for the given question."""
    sentences = [s.strip() for s in re.split(r'(?<=[.!?])\s+', chunk_text) if len(s.strip()) > 10]
    if len(sentences) <= max_sentences:
        return chunk_text

    q_emb = model.encode([question])
    s_embs = model.encode(sentences)
    sims = cosine_similarity(q_emb, s_embs)[0]

    # Also boost by keyword overlap
    for i, sent in enumerate(sentences):
        sims[i] = sims[i] * 0.7 + _keyword_boost(question, sent) * 0.3

    top_idx = np.argsort(sims)[::-1][:max_sentences]
    # Return sentences in their original order for readability
    top_idx_sorted = sorted(top_idx)
    return " ".join(sentences[i] for i in top_idx_sorted)


@app.post("/documents/{document_id}/ask", tags=["Document Q&A"])
def ask_document_question(
    document_id: int,
    req: QuestionRequest,
    db: Session = Depends(get_db),
    current_user: DBUser = Depends(get_current_user)
):
    """Ask a question about a document — uses hybrid semantic + keyword search for accurate answers."""
    doc = db.query(DBDocument).filter(
        DBDocument.id == document_id,
        DBDocument.user_id == current_user.id
    ).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")

    chunks = doc.chunks or []
    if not chunks:
        raise HTTPException(status_code=400, detail="No text chunks available for this document.")

    model = vector_store.model

    # 1. Semantic similarity scores
    question_embedding = model.encode([req.question])
    chunk_embeddings = model.encode(chunks)
    semantic_scores = cosine_similarity(question_embedding, chunk_embeddings)[0]

    # 2. Keyword overlap scores
    keyword_scores = np.array([_keyword_boost(req.question, c) for c in chunks])

    # 3. Hybrid score: 60% semantic + 40% keyword
    hybrid_scores = semantic_scores * 0.6 + keyword_scores * 0.4

    # 4. Get top 5 chunks, filter by minimum threshold
    top_k = min(5, len(chunks))
    top_indices = np.argsort(hybrid_scores)[::-1][:top_k]

    relevant_chunks = []
    for idx in top_indices:
        if hybrid_scores[idx] > 0.08:
            # Extract only the most relevant sentences from each chunk
            best_text = _extract_best_sentences(req.question, chunks[idx], model)
            relevant_chunks.append({
                "chunk_index": int(idx),
                "text": best_text,
                "full_text": chunks[idx],
                "relevance_score": round(float(hybrid_scores[idx]), 4),
                "semantic_score": round(float(semantic_scores[idx]), 4),
                "keyword_score": round(float(keyword_scores[idx]), 4)
            })

    # 5. Build a structured answer
    if relevant_chunks:
        # Combine best sentences from top chunks, deduplicate
        seen = set()
        answer_parts = []
        for c in relevant_chunks[:3]:
            for sent in re.split(r'(?<=[.!?])\s+', c["text"]):
                sent = sent.strip()
                if sent and sent not in seen and len(sent) > 10:
                    seen.add(sent)
                    answer_parts.append(sent)

        if answer_parts:
            answer = "Based on the document:\n\n" + " ".join(answer_parts)
        else:
            answer = "Based on the document:\n\n" + relevant_chunks[0]["text"]

        confidence = "high" if relevant_chunks[0]["relevance_score"] > 0.4 else "medium" if relevant_chunks[0]["relevance_score"] > 0.2 else "low"
    else:
        answer = "I couldn't find content in the document that's directly relevant to your question. Try rephrasing or asking about a different topic covered in the document."
        confidence = "none"

    return {
        "status": "success",
        "question": req.question,
        "answer": answer,
        "confidence": confidence,
        "relevant_chunks": relevant_chunks,
        "total_chunks_searched": len(chunks)
    }


@app.post("/documents/{document_id}/generate-qa", tags=["Document Q&A"])
def generate_document_qa(
    document_id: int,
    db: Session = Depends(get_db),
    current_user: DBUser = Depends(get_current_user)
):
    """Auto-generate questions and answers from the document content."""
    doc = db.query(DBDocument).filter(
        DBDocument.id == document_id,
        DBDocument.user_id == current_user.id
    ).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")

    chunks = doc.chunks or []
    if not chunks:
        raise HTTPException(status_code=400, detail="No text chunks available.")

    qa_pairs = []

    for i, chunk in enumerate(chunks[:10]):  # limit to first 10 chunks
        sentences = [s.strip() for s in re.split(r'(?<=[.!?])\s+', chunk) if len(s.strip()) > 30]
        if not sentences:
            continue

        # Generate different question types based on content
        main_sentence = sentences[0]

        # What question
        qa_pairs.append({
            "id": len(qa_pairs) + 1,
            "question": f"What does the document say about: \"{main_sentence[:80]}...\"?",
            "answer": chunk,
            "chunk_index": i,
            "type": "content"
        })

        # If chunk has multiple sentences, create a summary question
        if len(sentences) >= 2:
            qa_pairs.append({
                "id": len(qa_pairs) + 1,
                "question": f"Can you summarize the key points from section {i+1} of the document?",
                "answer": chunk,
                "chunk_index": i,
                "type": "summary"
            })

        # Detail question from a specific sentence
        if len(sentences) >= 3:
            detail_sent = sentences[1]
            qa_pairs.append({
                "id": len(qa_pairs) + 1,
                "question": f"Explain in detail: \"{detail_sent[:80]}...\"",
                "answer": detail_sent,
                "chunk_index": i,
                "type": "detail"
            })

    return {
        "status": "success",
        "document_id": document_id,
        "filename": doc.filename,
        "qa_pairs": qa_pairs[:15],  # limit to 15 Q&A pairs
        "total_generated": min(len(qa_pairs), 15)
    }


@app.delete("/documents/{document_id}", tags=["Document Q&A"])
def delete_document(
    document_id: int,
    db: Session = Depends(get_db),
    current_user: DBUser = Depends(get_current_user)
):
    """Delete an uploaded document."""
    doc = db.query(DBDocument).filter(
        DBDocument.id == document_id,
        DBDocument.user_id == current_user.id
    ).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")
    db.delete(doc)
    db.commit()
    return {"status": "success", "message": f"Document '{doc.filename}' deleted."}